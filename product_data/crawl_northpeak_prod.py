import os
from datetime import datetime
import tkinter as tk
from tkinter import filedialog

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException


def ask_save_path(default_filename: str) -> str:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    path = filedialog.asksaveasfilename(
        title="엑셀 파일 저장 위치를 선택해주세요",
        defaultextension=".xlsx",
        initialfile=default_filename,
        filetypes=[("Excel Workbook", "*.xlsx")],
    )

    root.destroy()
    return path


def create_driver() -> webdriver.Chrome:
    options = Options()
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-background-networking")
    options.add_argument("--disable-sync")
    options.add_argument("--no-default-browser-check")
    options.add_argument("--no-first-run")
    options.add_argument("--log-level=3")
    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(30)
    return driver


def close_popups(driver: webdriver.Chrome) -> None:
    # 1) JS alert(알림창) 닫기
    try:
        WebDriverWait(driver, 2).until(EC.alert_is_present())
        driver.switch_to.alert.accept()
    except Exception:
        pass

    # 2) 화면 레이어 팝업 닫기(있으면 누르고, 없으면 패스)
    # 사이트마다 버튼/문구가 달라서 "닫기" 계열 텍스트를 넓게 커버합니다.
    close_xpaths = [
        '//a[normalize-space()="닫기"]',
        '//button[normalize-space()="닫기"]',
        '//span[normalize-space()="닫기"]',
        '//*[contains(normalize-space(),"오늘 하루 보지 않기")]',
        '//*[contains(@class,"close")]',
        '//*[@id="layerClose" or @id="popupClose"]',
    ]

    for xp in close_xpaths:
        try:
            els = driver.find_elements(By.XPATH, xp)
            for el in els:
                if el.is_displayed() and el.is_enabled():
                    el.click()
                    return
        except Exception:
            continue


def crawl_product_names(driver: webdriver.Chrome) -> list[str]:
    url = "https://northpeak.co.kr/goods/goods_list.php?cateCd=005"
    driver.get(url)

    close_popups(driver)

    wait = WebDriverWait(driver, 15)

    # 핵심: strong가 아니라 "상품 상세 링크(a)"를 잡습니다.
    # ul/li 안에 있으면서 href에 goods_view.php가 포함된 링크만 대상으로 함
    product_link_xpath = '//ul/li//a[contains(@href,"goods_view.php") and normalize-space()]'

    try:
        wait.until(EC.presence_of_all_elements_located((By.XPATH, product_link_xpath)))
    except TimeoutException:
        print("요소 로딩 시간 초과(Timeout).")
        print("현재 URL:", driver.current_url)
        print("현재 title:", driver.title)
        raise

    elements = driver.find_elements(By.XPATH, product_link_xpath)

    names: list[str] = []
    for el in elements:
        t = (el.text or "").strip()
        if t:
            names.append(t)

    # 링크가 여러 군데(관련상품/최근본상품 등)에도 있을 수 있어 중복 제거(순서 유지)
    seen = set()
    unique: list[str] = []
    for n in names:
        if n not in seen:
            seen.add(n)
            unique.append(n)

    return unique


def save_to_excel(path: str, names: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "northpeak"

    ws["A1"] = "no"
    ws["B1"] = "상품명"

    for idx, name in enumerate(names, start=1):
        ws.cell(row=idx + 1, column=1, value=idx)
        ws.cell(row=idx + 1, column=2, value=name)

    ws.column_dimensions[get_column_letter(1)].width = 6
    ws.column_dimensions[get_column_letter(2)].width = 50

    dirpath = os.path.dirname(path)
    if dirpath:
        os.makedirs(dirpath, exist_ok=True)

    wb.save(path)


def main():
    default_filename = f"northpeak_names_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    save_path = ask_save_path(default_filename)

    if not save_path:
        print("저장을 취소하셨습니다.")
        return

    driver = create_driver()
    try:
        names = crawl_product_names(driver)
    except Exception as e:
        print("실패 원인:", repr(e))
        return
    finally:
        driver.quit()

    if not names:
        print("수집된 상품명이 없습니다. 선택자 또는 로딩을 확인해주세요.")
        return

    try:
        save_to_excel(save_path, names)
    except Exception as e:
        print("엑셀 저장 실패:", repr(e))
        return

    print(f"저장 완료: {save_path}")
    print(f"수집 건수: {len(names)}")


if __name__ == "__main__":
    main()
