
from pathlib import Path as _ConfigPath
import sys as _config_sys
_config_sys.path.insert(0, str(_ConfigPath(__file__).resolve().parents[1]))
from commerce_config import configured_url, output_dir
_GODOMALL_BASE_URL = configured_url('godomall_base_url')

import time
from datetime import datetime
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

import tkinter as tk
from tkinter import messagebox, filedialog

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


BASE_LIST_URL = _GODOMALL_BASE_URL + '/goods/goods_list.php?detailSearch=&delFl=n&scmFl=all&key=goodsNm&keyword=&searchDateFl=regDt&searchDate%5B%5D=&searchDate%5B%5D=&searchDate%5B%5D=&searchPeriod=-1&cateGoods%5B%5D=&cateGoods%5B%5D=&cateGoods%5B%5D=&cateGoods%5B%5D=&displayTheme%5B%5D=&displayTheme%5B%5D=&goodsPrice%5B%5D=&goodsPrice%5B%5D=&mileage%5B%5D=&mileage%5B%5D=&mileageFl=&stock%5B%5D=&stock%5B%5D=&optionFl=&optionTextFl=&addGoodsFl=&goodsDisplayFl=&goodsSellFl=&goodsDisplayMobileFl=&goodsSellMobileFl=&stockFl=&soldOut=&goodsIconCdPeriod=&goodsIconCd=&goodsDeliveryFl=&goodsDeliveryFixFl%5B%5D=all&event_text=&eventThemeSno=&sort=g.goodsNo+desc&pageNum=500&searchFl=y&applyPath=%2Fgoods%2Fgoods_list.php'

LIST_TBODY_XPATH = "/html/body/div[1]/div/div[2]/div[2]/div/form[2]/div[2]/table/tbody"

ROW_NAME_XPATH = "/html/body/div[1]/div/div[2]/div[2]/div/form[2]/div[2]/table/tbody/tr[{i}]/td[5]/div[1]/a"
ROW_LINK_XPATH = "/html/body/div[1]/div/div[2]/div[2]/div/form[2]/div[2]/table/tbody/tr[{i}]/td[4]/a"


def make_driver():
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    return webdriver.Chrome(options=options)


def _tk_root():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    return root


def show_login_required_popup():
    root = _tk_root()
    messagebox.showinfo(
        "로그인 필요",
        "관리자 페이지 로그인이 필요합니다.\n\n브라우저에서 로그인 후 확인을 눌러주세요."
    )
    root.destroy()


def ask_user_to_search_popup():
    root = _tk_root()
    messagebox.showinfo(
        "검색 진행",
        "원하시는 조건으로 상품 검색을 진행하신 뒤,\n검색 결과가 나온 상태에서 확인을 눌러주세요."
    )
    root.destroy()


def ask_continue_new_search() -> bool:
    root = _tk_root()
    cont = messagebox.askyesno(
        "완료",
        "엑셀 저장이 완료되었습니다.\n\n다른 검색 조건으로 계속 진행하시겠습니까?\n(예: 예=계속 / 아니오=종료)"
    )
    root.destroy()
    return cont


def ask_save_path(default_filename: str) -> str:
    """
    사용자가 저장할 위치/파일명을 선택.
    취소하면 "" 반환.
    """
    root = _tk_root()
    path = filedialog.asksaveasfilename(
        title="엑셀 저장 위치 선택",
        defaultextension=".xlsx",
        initialfile=default_filename,
        filetypes=[("Excel Workbook", "*.xlsx")]
    )
    root.destroy()
    return path or ""


def build_page_url(base_url: str, page: int) -> str:
    parts = urlsplit(base_url)
    q = parse_qsl(parts.query, keep_blank_values=True)
    q = [(k, v) for (k, v) in q if k.lower() != "page"]

    if page >= 2:
        q.insert(0, ("page", str(page)))

    new_query = urlencode(q, doseq=True)
    return urlunsplit((parts.scheme, parts.netloc, parts.path, new_query, parts.fragment))


def wait_until_list_ready(driver, timeout=60) -> bool:
    wait = WebDriverWait(driver, timeout)
    try:
        wait.until(EC.presence_of_element_located((By.XPATH, LIST_TBODY_XPATH)))
        return True
    except TimeoutException:
        return False


def list_has_first_row(driver) -> bool:
    try:
        driver.find_element(By.XPATH, ROW_LINK_XPATH.format(i=1))
        return True
    except NoSuchElementException:
        return False


def open_new_tab_and_get_url(driver, row_i: int, timeout=20) -> str:
    main_handle = driver.current_window_handle
    a = driver.find_element(By.XPATH, ROW_LINK_XPATH.format(i=row_i))

    before = set(driver.window_handles)
    a.click()

    try:
        WebDriverWait(driver, 10).until(lambda d: len(d.window_handles) > len(before))
        after = set(driver.window_handles)
        new_handle = list(after - before)[0]

        driver.switch_to.window(new_handle)
        WebDriverWait(driver, timeout).until(lambda d: d.current_url and d.current_url != "about:blank")
        time.sleep(0.3)

        opened_url = driver.current_url

        driver.close()
        driver.switch_to.window(main_handle)
        return opened_url

    except TimeoutException:
        WebDriverWait(driver, timeout).until(lambda d: d.current_url and d.current_url != "about:blank")
        time.sleep(0.3)
        opened_url = driver.current_url

        driver.back()
        wait_until_list_ready(driver, timeout=60)
        return opened_url


def extract_rows_from_current_page(driver) -> list:
    rows = []
    i = 1

    while True:
        try:
            name_el = driver.find_element(By.XPATH, ROW_NAME_XPATH.format(i=i))
            name = (name_el.text or "").strip()
            if not name:
                name = (name_el.get_attribute("textContent") or "").strip()
        except NoSuchElementException:
            break

        try:
            opened_url = open_new_tab_and_get_url(driver, i)
        except Exception as e:
            print(f"[ERROR] row={i} 새 탭 URL 추출 실패: {e}")
            opened_url = ""

        rows.append({
            "상품명": name,
            "링크주소": opened_url,
            "카테고리": "",
        })

        i += 2

    return rows


def save_to_excel(rows: list, output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "goods_list"

    headers = ["상품명", "링크주소", "카테고리"]
    ws.append(headers)

    for r in rows:
        ws.append([r.get("상품명", ""), r.get("링크주소", ""), r.get("카테고리", "")])

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 40 if col_idx == 2 else 30

    wb.save(output_path)


def crawl_by_base_url(driver, base_url: str) -> list:
    all_rows = []
    page = 1
    consecutive_empty_pages = 0

    while True:
        page_url = build_page_url(base_url, page)
        driver.get(page_url)

        if not wait_until_list_ready(driver, timeout=60):
            print(f"[STOP] page={page} 목록 로딩 실패")
            break

        if not list_has_first_row(driver):
            consecutive_empty_pages += 1
            print(f"[EMPTY] page={page} 연속 빈페이지={consecutive_empty_pages}")

            if consecutive_empty_pages >= 3:
                break

            page += 1
            continue

        consecutive_empty_pages = 0

        rows = extract_rows_from_current_page(driver)
        print(f"[PAGE] page={page} collected={len(rows)}")
        all_rows.extend(rows)

        if len(rows) == 0:
            print(f"[WARN] page={page} collected=0 (DOM 구조 변경 가능). 종료합니다.")
            break

        page += 1

    return all_rows


def run():
    driver = make_driver()
    driver.get(BASE_LIST_URL)

    show_login_required_popup()

    if not wait_until_list_ready(driver, timeout=300):
        print("로그인/목록 로딩 확인 실패(300초). 종료합니다.")
        driver.quit()
        return

    while True:
        ask_user_to_search_popup()

        base_url = driver.current_url
        print(f"[BASE_URL] {base_url}")

        rows = crawl_by_base_url(driver, base_url)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_filename = f"goods_list_{ts}.xlsx"
        save_path = ask_save_path(default_filename)

        if not save_path:
            root = _tk_root()
            messagebox.showwarning("취소", "저장이 취소되어 이번 결과는 저장하지 않습니다.")
            root.destroy()
        else:
            save_to_excel(rows, save_path)
            print(f"[SAVED] {save_path} (rows={len(rows)})")

        if not ask_continue_new_search():
            break

    driver.quit()


if __name__ == "__main__":
    run()
