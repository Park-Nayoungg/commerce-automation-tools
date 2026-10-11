
from pathlib import Path as _ConfigPath
import sys as _config_sys
_config_sys.path.insert(0, str(_ConfigPath(__file__).resolve().parents[1]))
from commerce_config import configured_url, output_dir
_GODOMALL_BASE_URL = configured_url('godomall_base_url')

import re
import time
import os
import threading
import unicodedata
from datetime import datetime
from tkinter import Tk, messagebox, Button, filedialog

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException, InvalidSessionIdException, NoSuchWindowException


stop_flag = False
driver = None

LOGIN_URL = _GODOMALL_BASE_URL + '/base/login.php'
GOODS_LIST_URL = _GODOMALL_BASE_URL + '/goods/goods_list.php'

OPTION_REGISTER_BTN_XPATH = "/html/body/div[1]/div/form/div[41]/table/tbody[2]/tr/td/button"
OPTION_SAVE_BTN_XPATH = "/html/body/div[1]/div/form/div[1]/div/input"
PRODUCT_SAVE_BTN_XPATH = "/html/body/div[1]/div/form/div[86]/div[3]/input"
NO_OPTION_STOCK_INPUT_XPATH = "/html/body/div[1]/div/form/div[29]/table/tbody/tr[3]/td[2]/label/input"
NO_OPTION_STOCK_INPUT_CSS = "input[name='stockCnt']"

OPTION_NAME_INPUT_CSS = "input[id='option_optionName_layer_0']"
OPTION_VALUE_TABLE_CSS = "table#optionValueLayer0 tbody"
OPTION_APPLY_BTN_CSS = "input#optionApplyBtn"
OPTION_ADD_BTN_XPATH = "//input[@value='추가' and contains(@onclick, 'option_value_conf_add_layer(0)')]"

SEARCH_FORM_CSS = "form#frmSearchGoods"
SEARCH_KEY_SELECT_CSS = "select[name='key']"
SEARCH_KEYWORD_INPUT_CSS = "input[name='keyword']"
SEARCH_SUBMIT_CSS = "form#frmSearchGoods input[type='submit'][value='검색']"
SEARCH_RESULT_TABLE_CSS = "form#frmList table.table.table-rows tbody"

USE_HEADLESS = False
WAIT_SECONDS = 30
MIN_PRODUCT_SCORE = 95
MIN_OPTION_SCORE = 95

ACTION_DELAY = 1.5
SEARCH_DELAY = 2.5

log_wb = Workbook()
log_ws = log_wb.active
log_ws.title = "log"
log_ws.append([
    "시간",
    "상품명(엑셀)",
    "상품명(쇼핑몰)",
    "작업결과",
    "메시지"
])


def write_log(excel_name, mall_name, result, message=""):
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    log_ws.append([now, excel_name, mall_name, result, message])


def stop_program():
    global stop_flag
    stop_flag = True

    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    messagebox.showinfo("알림", "작업이 중간에 종료됩니다.")
    root.destroy()


def create_control_window():
    root = Tk()
    root.title("작업 제어")
    root.geometry("250x100")
    root.attributes("-topmost", True)
    btn = Button(root, text="▶ 작업 중간 종료", command=stop_program, fg="red")
    btn.pack(pady=20)
    root.mainloop()


def show_login_popup():
    def on_confirm():
        print("사용자가 로그인 완료를 확인했습니다.")
        root.destroy()

    root = Tk()
    root.title("로그인 요청")
    root.geometry("300x100")
    root.resizable(False, False)
    root.attributes("-topmost", True)
    Button(root, text="로그인을 완료한 후 이 버튼을 눌러주세요.", command=on_confirm).pack(pady=30)
    root.mainloop()


def ask_excel_file_path():
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    file_path = filedialog.askopenfilename(
        title="엑셀 파일 선택",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
    )
    root.destroy()

    if not file_path:
        raise Exception("엑셀 파일이 선택되지 않았습니다.")

    print(f"선택된 엑셀 파일: {file_path}")
    return file_path


def ask_driver_quit_popup():
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    result = messagebox.askyesno("종료 확인", "드라이버를 종료하시겠습니까?")
    root.destroy()
    return bool(result)


def safe_click(driver, element):
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", element)
    time.sleep(0.2)
    try:
        element.click()
    except Exception:
        driver.execute_script("arguments[0].click();", element)

def get_alive_window_handles(driver):
    try:
        return driver.window_handles[:]
    except (NoSuchWindowException, InvalidSessionIdException):
        return []

def is_window_alive(driver, handle):
    if not handle:
        return False
    try:
        return handle in driver.window_handles
    except (NoSuchWindowException, InvalidSessionIdException):
        return False

def safe_switch_window(driver, handle, timeout=5):
    end_at = time.time() + timeout

    while time.time() < end_at:
        try:
            if handle in driver.window_handles:
                driver.switch_to.window(handle)
                return True
        except Exception:
            pass
        time.sleep(0.2)

    return False

def close_window_if_alive(driver, handle):
    try:
        if is_window_alive(driver, handle):
            driver.switch_to.window(handle)
            driver.close()
            return True
    except Exception:
        pass
    return False

def cleanup_extra_windows(driver, keep_handles):
    keep_set = set(keep_handles)
    handles = get_alive_window_handles(driver)

    for handle in handles:
        if handle in keep_set:
            continue
        try:
            driver.switch_to.window(handle)
            driver.close()
        except Exception:
            pass

    for keep_handle in keep_handles:
        if safe_switch_window(driver, keep_handle, timeout=2):
            break

def wait_new_window_and_switch(driver, old_handles, timeout=30):
    WebDriverWait(driver, timeout).until(
        lambda d: len(d.window_handles) > len(old_handles)
    )

    new_handles = [h for h in driver.window_handles if h not in old_handles]
    if not new_handles:
        raise TimeoutException(
            f"새 창이 감지되었지만 새 handle 을 찾지 못했습니다. old={old_handles}, current={driver.window_handles}"
        )

    new_handle = new_handles[0]
    driver.switch_to.window(new_handle)
    return new_handle

def get_option_rows(driver, timeout=15):
    end_at = time.time() + timeout
    last_error = None

    while time.time() < end_at:
        try:
            table = driver.find_element(By.ID, "optionGridTable")
            tbody = table.find_element(By.TAG_NAME, "tbody")
            rows = tbody.find_elements(By.TAG_NAME, "tr")

            data_rows = []
            for row in rows:
                row_id = (row.get_attribute("id") or "").strip()
                if row_id.startswith("tbl_option_info_"):
                    data_rows.append(row)

            if data_rows or rows:
                return data_rows
        except Exception as e:
            last_error = e

        time.sleep(0.5)

    if last_error:
        raise last_error
    return []

def wait_option_popup_ready(driver, timeout=40):
    WebDriverWait(driver, timeout).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, OPTION_NAME_INPUT_CSS))
    )
    WebDriverWait(driver, timeout).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, OPTION_VALUE_TABLE_CSS))
    )
    WebDriverWait(driver, timeout).until(
        lambda d: len(d.find_elements(By.ID, "optionGridTable")) > 0
    )
    WebDriverWait(driver, timeout).until(
        lambda d: len(get_option_rows(d)) >= 0
    )
    time.sleep(1.5)

def wait_product_detail_ready(driver, timeout=30):
    selectors = [
        "input[name='goodsNm']",
        "input[name='goodsName']",
        "#goodsNm",
        "#goodsName",
    ]

    WebDriverWait(driver, timeout).until(
        lambda d: any(
            len(d.find_elements(By.CSS_SELECTOR, selector)) > 0
            for selector in selectors
        )
    )

def wait_option_popup_opened(driver, old_handles, timeout=30):
    original_handle = None
    try:
        original_handle = driver.current_window_handle
    except Exception:
        pass

    def condition(d):
        current_handles = d.window_handles[:]
        new_handles = [h for h in current_handles if h not in old_handles]
        if new_handles:
            return ("new_window", new_handles[0])

        for handle in current_handles:
            try:
                d.switch_to.window(handle)
                if d.find_elements(By.CSS_SELECTOR, OPTION_NAME_INPUT_CSS):
                    return ("existing_window", handle)
            except Exception:
                continue

        return False

    result = WebDriverWait(driver, timeout).until(condition)
    mode, handle = result
    driver.switch_to.window(handle)
    return mode, handle

def normalize_text(value: str) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text).upper().strip()
    text = text.replace("&", " AND ")
    text = re.sub(r"\[[^\]]*\]", " ", text)
    text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"[-_/,+.]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def make_search_keyword(product_name: str) -> str:
    text = "" if product_name is None else str(product_name)
    text = unicodedata.normalize("NFKC", text).upper().strip()
    text = re.sub(r"[^0-9A-Z가-힣]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_option_text(value: str) -> str:
    text = normalize_text(value)
    text = re.sub(r"\b(재고|품절|OPTION|옵션|COLOR|COLOUR|SIZE)\b", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"재고\s*\(\s*\d+\s*\)", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def match_score(a: str, b: str) -> int:
    if not a or not b:
        return 0
    if a == b:
        return 100
    if a in b or b in a:
        return 90

    a_tokens = set(a.split())
    b_tokens = set(b.split())
    if not a_tokens or not b_tokens:
        return 0

    common = a_tokens & b_tokens
    union = a_tokens | b_tokens
    return int((len(common) / len(union)) * 100)


def parse_int(value) -> int:
    digits = re.sub(r"[^\d]", "", str(value or ""))
    return int(digits) if digits else 0


def calc_target_stock_for_soldout(current_stock: int) -> int:
    if current_stock >= 100:
        return current_stock - 100
    return current_stock

def calc_target_stock_for_available(current_stock: int) -> int:
    if current_stock == 0:
        return 100
    return current_stock

def parse_option_line(line: str):
    raw = str(line or "").strip()
    if not raw:
        return None

    sold_out = "[품절]" in raw
    stock_match = re.search(r"재고\((\d+)\)", raw)
    excel_stock = int(stock_match.group(1)) if stock_match else 0

    option_name = raw
    option_name = option_name.replace("[품절]", " ")
    option_name = re.sub(r"재고\(\d+\)", " ", option_name)
    option_name = re.sub(r"\[[^\]]+\]", " ", option_name)
    option_name = re.sub(r"\s+", " ", option_name).strip()

    if not option_name:
        return None

    return {
        "raw": raw,
        "name": option_name,
        "norm": normalize_option_text(option_name),
        "sold_out": sold_out,
        "excel_stock": excel_stock,
    }


def load_excel_context(excel_path: str):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[wb.sheetnames[0]]

    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    header_index = {name: idx for idx, name in enumerate(headers)}

    required_headers = ["productName", "optionFl", "optionSummary", "stockCnt"]
    for header in required_headers:
        if header not in header_index:
            raise Exception(f"엑셀에 '{header}' 컬럼이 없습니다.")

    products = []
    stock_map = {}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        product_name = str(row[header_index["productName"]] or "").strip()
        option_fl = str(row[header_index["optionFl"]] or "").strip().lower()
        option_summary = str(row[header_index["optionSummary"]] or "").strip()
        stock_cnt = parse_int(row[header_index["stockCnt"]])

        if not product_name:
            continue

        parsed_options = []
        if option_fl == "y" and option_summary:
            for line in option_summary.splitlines():
                parsed = parse_option_line(line)
                if parsed:
                    parsed_options.append(parsed)

        product = {
            "row_idx": row_idx,
            "product_name": product_name,
            "optionFl": option_fl,
            "options": parsed_options,
            "stockCnt": stock_cnt,
        }
        products.append(product)

        key = normalize_text(product_name)
        if key not in stock_map:
            stock_map[key] = product

    return {
        "path": excel_path,
        "wb": wb,
        "ws": ws,
        "header_index": header_index,
        "products": products,
        "stock_map": stock_map,
    }

def mark_product_name_red(excel_ctx, excel_product):
    ws = excel_ctx["ws"]
    header_index = excel_ctx["header_index"]
    row_idx = excel_product["row_idx"]
    col_idx = header_index["productName"] + 1
    cell = ws.cell(row=row_idx, column=col_idx)
    cell.font = Font(color="FF0000")

def mark_option_summary_red(excel_ctx, excel_product):
    ws = excel_ctx["ws"]
    header_index = excel_ctx["header_index"]
    row_idx = excel_product["row_idx"]
    col_idx = header_index["optionSummary"] + 1
    cell = ws.cell(row=row_idx, column=col_idx)
    cell.font = Font(color="FF0000")

def mark_product_name_green(excel_ctx, excel_product):
    ws = excel_ctx["ws"]
    header_index = excel_ctx["header_index"]
    row_idx = excel_product["row_idx"]
    col_idx = header_index["productName"] + 1
    cell = ws.cell(row=row_idx, column=col_idx)
    cell.font = Font(color="008000")

def mark_option_summary_green(excel_ctx, excel_product):
    ws = excel_ctx["ws"]
    header_index = excel_ctx["header_index"]
    row_idx = excel_product["row_idx"]
    col_idx = header_index["optionSummary"] + 1
    cell = ws.cell(row=row_idx, column=col_idx)
    cell.font = Font(color="008000")

def clean_option_display_name(raw_name: str) -> str:
    text = "" if raw_name is None else str(raw_name)
    text = text.replace("[품절]", " ")
    text = re.sub(r"\[[^\]]+\]", " ", text)
    text = re.sub(r"재고\(\d+\)", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def save_excel_with_marks(excel_ctx):
    original_path = excel_ctx["path"]
    save_path = original_path[:-5] + "_result.xlsx"
    if original_path.lower().endswith(".xlsx"):
        save_path = original_path[:-5] + "_result.xlsx"
    else:
        save_path = original_path + "_result.xlsx"

    excel_ctx["wb"].save(save_path)
    print(f"표시 반영 엑셀 저장 완료: {save_path}")
    return save_path

def get_excel_product_list(excel_ctx):
    return excel_ctx["products"]

def is_site_option_product(driver):
    try:
        btns = driver.find_elements(By.XPATH, OPTION_REGISTER_BTN_XPATH)
        if not btns:
            return False

        btn = btns[0]
        text = (btn.text or "").strip()
        return "옵션" in text
    except Exception:
        return False

def get_product_name_in_detail(driver):
    candidates = [
        "input[name='goodsNm']",
        "input[name='goodsName']",
        "#goodsNm",
        "#goodsName",
    ]

    for selector in candidates:
        try:
            element = driver.find_element(By.CSS_SELECTOR, selector)
            value = (element.get_attribute("value") or "").strip()
            if value:
                return value
        except Exception:
            pass

    try:
        title = driver.title.strip()
        if title:
            return title
    except Exception:
        pass

    return ""

def run_goods_search(driver, wait, product_name: str):
    keyword = make_search_keyword(product_name)

    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, SEARCH_FORM_CSS)))

    try:
        key_select = driver.find_element(By.CSS_SELECTOR, SEARCH_KEY_SELECT_CSS)
        driver.execute_script("arguments[0].value = 'goodsNm';", key_select)
        driver.execute_script("arguments[0].dispatchEvent(new Event('change', { bubbles: true }));", key_select)
    except Exception:
        pass

    keyword_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, SEARCH_KEYWORD_INPUT_CSS)))
    keyword_input.click()
    keyword_input.clear()
    keyword_input.send_keys(keyword)

    submit_btn = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, SEARCH_SUBMIT_CSS)))
    safe_click(driver, submit_btn)

    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, SEARCH_RESULT_TABLE_CSS)))
    time.sleep(SEARCH_DELAY)

    print(f"[검색 실행] 원본='{product_name}' / 검색어='{keyword}'")


def get_search_result_main_rows(driver):
    tbody = driver.find_element(By.CSS_SELECTOR, SEARCH_RESULT_TABLE_CSS)
    rows = tbody.find_elements(By.TAG_NAME, "tr")

    main_rows = []
    for row in rows:
        try:
            goods_name_el = row.find_element(By.CSS_SELECTOR, "td:nth-child(4) a.text-blue.hand")
            if goods_name_el:
                main_rows.append(row)
        except Exception:
            continue

    return main_rows


def get_row_goods_name(row):
    try:
        return row.find_element(By.CSS_SELECTOR, "td:nth-child(4) a.text-blue.hand").text.strip()
    except Exception:
        return ""


def get_row_goods_link(row):
    return row.find_element(By.CSS_SELECTOR, "td:nth-child(4) a.text-blue.hand")


def find_best_search_result_row(driver, excel_product):
    rows = get_search_result_main_rows(driver)
    if not rows:
        return None, -1

    best_row = None
    best_score = -1
    excel_norm = normalize_text(excel_product["product_name"])

    for row in rows:
        mall_name = get_row_goods_name(row)
        mall_norm = normalize_text(mall_name)
        score = match_score(excel_norm, mall_norm)

        if score > best_score:
            best_score = score
            best_row = row

    return best_row, best_score

def open_product_detail_from_search_row(driver, row):
    origin = driver.current_window_handle
    old_handles = driver.window_handles[:]

    goods_link = get_row_goods_link(row)

    driver.execute_script("window.focus();")

    safe_click(driver, goods_link)

    wait_new_window_and_switch(driver, old_handles, timeout=30)
    wait_product_detail_ready(driver, timeout=30)
    return origin

def open_option_popup(driver, wait):
    detail_handle = driver.current_window_handle
    old_handles = driver.window_handles[:]

    btn = wait.until(EC.element_to_be_clickable((By.XPATH, OPTION_REGISTER_BTN_XPATH)))

    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
    driver.execute_script("arguments[0].focus();", btn)
    driver.execute_script("window.focus();")
    time.sleep(1.0)

    try:
        btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", btn)

    time.sleep(1.0)

    mode, option_handle = wait_option_popup_opened(driver, old_handles, timeout=40)
    wait_option_popup_ready(driver, timeout=40)
    return detail_handle, option_handle

def switch_back_to_parent_window(driver, parent_handle, timeout=10):
    end_at = time.time() + timeout

    while time.time() < end_at:
        handles = driver.window_handles[:]

        if parent_handle in handles:
            try:
                driver.switch_to.window(parent_handle)
                return
            except Exception:
                pass

        time.sleep(0.2)

    raise Exception(f"부모 상세창 복귀 실패: parent_handle={parent_handle}, handles={driver.window_handles}")

def ensure_option_name_layer(driver, option_name="색상"):
    input_el = driver.find_element(By.CSS_SELECTOR, OPTION_NAME_INPUT_CSS)
    current = (input_el.get_attribute("value") or "").strip()
    if current != option_name:
        input_el.click()
        input_el.clear()
        input_el.send_keys(option_name)
        time.sleep(ACTION_DELAY)

def get_option_value_inputs(driver):
    tbody = driver.find_element(By.CSS_SELECTOR, OPTION_VALUE_TABLE_CSS)
    return tbody.find_elements(By.CSS_SELECTOR, "input[id^='option_optionValue_0_']")

def find_empty_option_value_input(driver):
    inputs = get_option_value_inputs(driver)
    for input_el in inputs:
        value = (input_el.get_attribute("value") or "").strip()
        if not value:
            return input_el
    return None

def add_option_value_row(driver, wait):
    add_buttons = driver.find_elements(By.XPATH, OPTION_ADD_BTN_XPATH)
    if not add_buttons:
        raise Exception("옵션값 추가 버튼을 찾지 못했습니다.")

    before_count = len(get_option_value_inputs(driver))
    safe_click(driver, add_buttons[0])

    end_at = time.time() + 5
    while time.time() < end_at:
        after_count = len(get_option_value_inputs(driver))
        if after_count > before_count:
            return
        time.sleep(0.2)

    raise Exception("옵션값 입력행이 추가되지 않았습니다.")

def append_option_value_to_layer(driver, wait, option_value: str):
    option_value = clean_option_display_name(option_value)

    input_el = find_empty_option_value_input(driver)
    if input_el is None:
        add_option_value_row(driver, wait)
        input_el = find_empty_option_value_input(driver)

    if input_el is None:
        raise Exception("새 옵션값 입력칸을 확보하지 못했습니다.")

    input_el.click()
    input_el.clear()
    input_el.send_keys(option_value)
    time.sleep(ACTION_DELAY)

def apply_option_grid(driver, wait):
    apply_btn = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, OPTION_APPLY_BTN_CSS)))
    safe_click(driver, apply_btn)
    time.sleep(ACTION_DELAY)

    try:
        alert = driver.switch_to.alert
        print(f"[옵션 재적용 알림] {alert.text}")
        alert.accept()
        time.sleep(1.5)
    except Exception:
        pass

def find_or_create_option_row(driver, wait, excel_product, excel_option):
    rows = get_option_rows(driver)
    row, score = find_best_option_row(rows, excel_option)
    if row is not None:
        return row, score, False

    if not excel_option["sold_out"]:
        return None, -1, False

    ensure_option_name_layer(driver, "색상")
    append_option_value_to_layer(driver, wait, excel_option["name"])
    apply_option_grid(driver, wait)

    rows = get_option_rows(driver)
    row, score = find_best_option_row(rows, excel_option)
    if row is not None:
        return row, score, True

    return None, -1, True




def get_row_option_name(row):
    try:
        hidden = row.find_element(By.CSS_SELECTOR, "input[name='optionY[optionValueText][]']")
        value = (hidden.get_attribute("value") or "").strip()
        if value:
            return value
    except Exception:
        pass

    try:
        td = row.find_element(By.CSS_SELECTOR, "td.colOptionValueLast")
        value = td.text.strip()
        if value:
            return value
    except Exception:
        pass

    return ""

def get_row_stock_input(row, timeout=10):
    end_at = time.time() + timeout
    last_error = None

    while time.time() < end_at:
        try:
            input_el = row.find_element(By.CSS_SELECTOR, "td.colStockCnt input[name='optionY[stockCnt][]']")
            if input_el.is_displayed():
                return input_el
            return input_el
        except Exception as e:
            last_error = e
            time.sleep(0.3)

    if last_error:
        raise last_error
    raise Exception("옵션 재고 입력칸을 찾지 못했습니다.")

def find_best_option_row(rows, excel_option):
    best_row = None
    best_score = -1

    for row in rows:
        mall_name = get_row_option_name(row)
        mall_norm = normalize_option_text(mall_name)
        score = match_score(excel_option["norm"], mall_norm)
        if score > best_score:
            best_score = score
            best_row = row

    if best_score < MIN_OPTION_SCORE:
        return None, best_score
    return best_row, best_score

def is_window_alive(driver, handle):
    try:
        return handle in driver.window_handles
    except Exception:
        return False

def safe_switch_window(driver, handle, timeout=5):
    end_at = time.time() + timeout

    while time.time() < end_at:
        try:
            if handle in driver.window_handles:
                driver.switch_to.window(handle)
                return True
        except Exception:
            pass
        time.sleep(0.2)

    return False

def save_option_popup(driver, wait, option_handle, detail_handle):
    save_btn = wait.until(EC.element_to_be_clickable((By.XPATH, OPTION_SAVE_BTN_XPATH)))
    safe_click(driver, save_btn)
    time.sleep(ACTION_DELAY)

    try:
        alert = driver.switch_to.alert
        print(f"[옵션 저장 알림] {alert.text}")
        alert.accept()
        time.sleep(1.5)
    except Exception:
        pass

    time.sleep(1.5)

    if not is_window_alive(driver, option_handle):
        print("[옵션 저장 후] 옵션창 자동 닫힘 감지")
    else:
        print("[옵션 저장 후] 옵션창이 아직 열려 있음, 직접 닫기 시도")
        try:
            driver.switch_to.window(option_handle)
            driver.close()
            time.sleep(0.5)
        except Exception:
            pass

    if not safe_switch_window(driver, detail_handle, timeout=5):
        raise Exception("옵션 저장 후 상세창 복귀 실패")


def get_no_option_stock_input(driver):
    try:
        return driver.find_element(By.XPATH, NO_OPTION_STOCK_INPUT_XPATH)
    except Exception:
        return driver.find_element(By.CSS_SELECTOR, NO_OPTION_STOCK_INPUT_CSS)


def update_input_value(driver, input_element, value: int):
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_element)
    time.sleep(ACTION_DELAY)

    try:
        input_element.click()
        input_element.clear()
        input_element.send_keys(str(value))
    except Exception:
        driver.execute_script("arguments[0].value = arguments[1];", input_element, str(value))
        driver.execute_script("arguments[0].dispatchEvent(new Event('input', { bubbles: true }));", input_element)
        driver.execute_script("arguments[0].dispatchEvent(new Event('change', { bubbles: true }));", input_element)


def save_product_detail(driver, wait):
    save_btn = wait.until(EC.element_to_be_clickable((By.XPATH, PRODUCT_SAVE_BTN_XPATH)))
    safe_click(driver, save_btn)
    time.sleep(2)

    try:
        alert = driver.switch_to.alert
        print(f"[상품 저장 알림] {alert.text}")
        alert.accept()
        time.sleep(ACTION_DELAY)
    except Exception:
        pass

def process_option_popup(driver, wait, excel_ctx, excel_product):
    rows = get_option_rows(driver)
    if not rows:
        print("[옵션] 옵션 행이 없습니다.")
        return "NO_CHANGE"

    changed = 0

    if not excel_product["options"]:
        print("[옵션] 엑셀 옵션 정보가 없습니다.")
        return "NO_CHANGE"

    for excel_option in excel_product["options"]:
        row, score, created = find_or_create_option_row(driver, wait, excel_product, excel_option)

        if row is None:
            print(f"[옵션 미매칭] {excel_option['raw']}")
            write_log(
                excel_product["product_name"],
                "",
                "옵션미매칭",
                f"옵션없음 / 추가실패 / {excel_option['raw']}"
            )
            mark_option_summary_red(excel_ctx, excel_product)
            continue

        if created:
            write_log(
                excel_product["product_name"],
                excel_option["name"],
                "옵션추가",
                f"품절옵션 자동추가 / {excel_option['raw']}"
            )

        stock_input = get_row_stock_input(row)
        current_stock = parse_int(stock_input.get_attribute("value"))

        if excel_option["sold_out"]:
            target_stock = calc_target_stock_for_soldout(current_stock)
        else:
            target_stock = calc_target_stock_for_available(current_stock)

        if current_stock == target_stock:
            print(f"[변경 없음] {excel_option['name']} / 현재={current_stock}")
            continue

        update_input_value(driver, stock_input, target_stock)
        updated_stock = parse_int(stock_input.get_attribute("value"))
        changed += 1

        mall_option_name = get_row_option_name(row)
        print(f"[옵션 변경] {mall_option_name} / {current_stock} -> {target_stock} / 반영값={updated_stock}")
        write_log(
            excel_product["product_name"],
            mall_option_name,
            "옵션변경",
            f"{current_stock} -> {target_stock}"
        )

    if changed == 0:
        return "NO_CHANGE"

    return "UPDATED"


def process_no_option_stock(driver, wait, excel_product):
    if excel_product.get("optionFl") == "y":
        raise Exception(f"엑셀 옵션상품인데 무옵션 처리 함수로 들어왔습니다: {excel_product['product_name']}")

    stock_input = get_no_option_stock_input(driver)
    current_stock = parse_int(stock_input.get_attribute("value"))
    target_stock = calc_target_stock_for_soldout(current_stock)

    if current_stock == 0:
        target_stock = 100
    else:
        target_stock = calc_target_stock_for_soldout(current_stock)

    if current_stock == target_stock:
        return "NO_CHANGE"

    update_input_value(driver, stock_input, target_stock)
    updated_stock = parse_int(stock_input.get_attribute("value"))

    print(f"[무옵션 변경] {current_stock} -> {target_stock} / 반영값={updated_stock}")

    save_product_detail(driver, wait)
    return "UPDATED"

def process_mismatch_option_as_no_option(driver, wait, excel_product):
    stock_input = get_no_option_stock_input(driver)
    current_stock = parse_int(stock_input.get_attribute("value"))

    if parse_int(excel_product.get("stockCnt", 0)) > 0:
        target_stock = 100
    else:
        target_stock = 0

    if current_stock == target_stock:
        return "NO_CHANGE"

    update_input_value(driver, stock_input, target_stock)
    updated_stock = parse_int(stock_input.get_attribute("value"))

    print(f"[옵션불일치-무옵션 처리] {current_stock} -> {target_stock} / 반영값={updated_stock}")

    save_product_detail(driver, wait)
    return "UPDATED"

def process_product_detail(driver, wait, excel_ctx, excel_product):
    mall_product_name = get_product_name_in_detail(driver)

    try:
        if not is_site_option_product(driver):

            if excel_product.get("optionFl") == "y":
                print(f"[옵션 불일치] 엑셀은 옵션상품인데 사이트는 무옵션: {excel_product['product_name']}")
                mark_product_name_green(excel_ctx, excel_product)
                mark_option_summary_green(excel_ctx, excel_product)

                mismatch_result = process_mismatch_option_as_no_option(driver, wait, excel_product)

                write_log(
                    excel_product["product_name"],
                    mall_product_name,
                    "옵션불일치",
                    f"엑셀 옵션상품 / 사이트 무옵션 / stockCnt={excel_product.get('stockCnt', 0)} / 결과={mismatch_result}"
                )

                if mismatch_result == "UPDATED":
                    print(f"[상품 저장 완료 - 옵션불일치 무옵션처리] {excel_product['product_name']}")

                return mismatch_result, mall_product_name

            no_option_result = process_no_option_stock(driver, wait, excel_product)

            if no_option_result == "UPDATED":
                print(f"[상품 저장 완료 - 무옵션상품] {excel_product['product_name']}")

            return no_option_result, mall_product_name

        try:
            option_parent, option_handle = open_option_popup(driver, wait)
        except Exception as e:
            print(f"[옵션 팝업 1차 실패] {excel_product['product_name']} / {e}")
            time.sleep(3.0)
            option_parent, option_handle = open_option_popup(driver, wait)
        option_result = process_option_popup(driver, wait, excel_ctx, excel_product)

        if option_result == "UPDATED":
            save_option_popup(
                driver,
                wait,
                option_handle=option_handle,
                detail_handle=option_parent
            )
        else:
            try:
                if option_handle in driver.window_handles:
                    driver.switch_to.window(option_handle)
                    driver.close()
            except Exception:
                pass

            safe_switch_window(driver, option_parent, timeout=5)

        if not safe_switch_window(driver, option_parent, timeout=5):
            raise Exception(f"옵션 저장 후 상세창 복귀 실패: {excel_product['product_name']}")

        if option_result != "UPDATED":
            return option_result, mall_product_name

        save_product_detail(driver, wait)
        print(f"[상품 저장 완료 - 옵션상품] {excel_product['product_name']}")
        return "UPDATED", mall_product_name

    except Exception as e:
        try:
            if "option_parent" in locals():
                switch_back_to_parent_window(driver, option_parent)
                cleanup_extra_windows(driver, keep_handles=[option_parent])
        except Exception:
            pass

        print(f"[상세 처리 오류] 상세상품명={mall_product_name} / 엑셀상품명={excel_product['product_name']} / {e}")
        raise

def process_excel_products(driver, wait, excel_ctx):
    global stop_flag

    products = get_excel_product_list(excel_ctx)
    total_count = len(products)

    driver.get(GOODS_LIST_URL)
    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, SEARCH_FORM_CSS)))

    for idx, excel_product in enumerate(products, start=1):
        if stop_flag:
            print("사용자 중단 요청으로 종료합니다.")
            return "STOPPED"

        try:
            driver.get(GOODS_LIST_URL)
            wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, SEARCH_FORM_CSS)))

            run_goods_search(driver, wait, excel_product["product_name"])

            row, score = find_best_search_result_row(driver, excel_product)
            if row is None:
                msg = f"검색 결과 없음 / 검색어={make_search_keyword(excel_product['product_name'])}"
                print(f"[검색 결과 없음] 상품명={excel_product['product_name']} / 검색어={make_search_keyword(excel_product['product_name'])}")
                write_log(excel_product["product_name"], "", "검색실패", msg)
                mark_product_name_red(excel_ctx, excel_product)
                continue

            if score < MIN_PRODUCT_SCORE:
                msg = f"유사도 낮음 / 점수={score} / 검색어={make_search_keyword(excel_product['product_name'])}"
                print(f"[유사도 낮음] 상품명={excel_product['product_name']} / 검색어={make_search_keyword(excel_product['product_name'])} / 점수={score}")
                write_log(excel_product["product_name"], "", "스킵", msg)
                continue

            origin = open_product_detail_from_search_row(driver, row)
            detail_result, mall_product_name = process_product_detail(driver, wait, excel_ctx, excel_product)

            print(f"[상품 결과] {excel_product['product_name']} / 점수={score} / 결과={detail_result}")
            write_log(excel_product["product_name"], mall_product_name, detail_result, f"유사도={score}")

            handles = driver.window_handles
            try:
                current_handle = driver.current_window_handle
            except Exception:
                current_handle = None

            try:
                current_handle = driver.current_window_handle
                if current_handle != origin:
                    driver.close()
                    time.sleep(0.5)
            except Exception:
                pass

            safe_switch_window(driver, origin, timeout=3)

            if origin in driver.window_handles:
                try:
                    driver.switch_to.window(origin)
                except Exception:
                    pass

            time.sleep(1)

        except InvalidSessionIdException as e:
            print(f"[세션 종료] {idx}/{total_count} / {excel_product['product_name']} / {e}")
            write_log(excel_product["product_name"], "", "세션종료", str(e))
            raise Exception(f"브라우저 세션이 종료되었습니다. 마지막 작업 상품: {excel_product['product_name']}")

        except NoSuchWindowException as e:
            print(f"[창 종료 감지] {idx}/{total_count} / {excel_product['product_name']} / {e}")
            write_log(excel_product["product_name"], "", "창종료", str(e))

            try:
                handles = get_alive_window_handles(driver)
                if handles:
                    base_handle = handles[0]
                    safe_switch_window(driver, base_handle, timeout=3)
            except Exception:
                pass

            continue
    
        except Exception as e:
            print(f"[처리 오류] {idx}/{total_count} / {excel_product['product_name']} / {e}")
            write_log(excel_product["product_name"], "", "에러", str(e))

            try:
                handles = driver.window_handles[:]
                if len(handles) >= 2:
                    base_handle = handles[0]
                    for handle in handles[1:]:
                        try:
                            driver.switch_to.window(handle)
                            driver.close()
                        except Exception:
                            pass
                    try:
                        driver.switch_to.window(base_handle)
                    except Exception:
                        pass
            except InvalidSessionIdException:
                return "SESSION_LOST"
            except Exception:
                pass

            continue

    return "DONE"


if __name__ == "__main__":
    excel_path = ask_excel_file_path()
    excel_ctx = load_excel_context(excel_path)

    options = webdriver.ChromeOptions()
    if USE_HEADLESS:
        options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--remote-debugging-port=9222")
    options.add_argument("--disable-popup-blocking")

    driver = webdriver.Chrome(options=options)
    wait = WebDriverWait(driver, WAIT_SECONDS)

    threading.Thread(target=create_control_window, daemon=True).start()

    try:
        driver.get(LOGIN_URL)
        show_login_popup()

        if stop_flag:
            root = Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            messagebox.showinfo("중단됨", "사용자 요청으로 작업이 중단되었습니다.\n브라우저는 그대로 유지됩니다.")
            root.destroy()
        else:
            result = process_excel_products(driver=driver, wait=wait, excel_ctx=excel_ctx)
            base_dir = os.path.dirname(excel_ctx["path"])

            result_excel_path = save_excel_with_marks(excel_ctx)
            log_path = os.path.join(
                base_dir,
                f"log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )
            log_wb.save(log_path)
            print(f"로그 엑셀 저장 완료: {log_path}")

            root = Tk()
            root.withdraw()
            root.attributes("-topmost", True)

            if result == "STOPPED":
                messagebox.showinfo("중단됨", f"작업이 중단되었습니다.\n결과 엑셀: {result_excel_path}\n로그 엑셀: {log_path}")
            elif result == "SESSION_LOST":
                messagebox.showerror("세션 종료", f"브라우저 세션이 종료되어 작업이 중단되었습니다.\n결과 엑셀: {result_excel_path}\n로그 엑셀: {log_path}")
            else:
                messagebox.showinfo("완료", f"엑셀 기준 상품 처리가 완료되었습니다.\n결과 엑셀: {result_excel_path}\n로그 엑셀: {log_path}")
            
            root.destroy()

        quit_driver = ask_driver_quit_popup()
        if quit_driver:
            driver.quit()

    except Exception as e:
        print(f"[치명오류] {e}")
        try:
            result_excel_path = save_excel_with_marks(excel_ctx)
        except Exception:
            result_excel_path = ""

        try:
            base_dir = os.path.dirname(excel_ctx["path"])
            log_path = os.path.join(
                base_dir,
                f"log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )
            log_wb.save(log_path)
        except Exception:
            log_path = ""

        try:
            root = Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            messagebox.showerror("오류", f"예기치 못한 오류가 발생했습니다.\n{e}\n결과 엑셀: {result_excel_path}\n로그 엑셀: {log_path}")
            root.destroy()
        except Exception:
            pass
