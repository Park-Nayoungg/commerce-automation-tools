import re
import time
import random
import threading
import queue
from datetime import datetime
import json

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from tkinter import filedialog
from urllib.parse import quote
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

from selenium import webdriver
import os

BASE = "http://parkersb2b.co.kr/"
SEARCH_URL = BASE + "goods/goods_search.php"
VIEW_URL = BASE + "goods/goods_view.php"
GOODS_PS_URL = BASE + "goods/goods_ps.php"

COOKIE_DOMAIN = "parkersb2b.co.kr"

SNAPSHOT_PATH = "parkers_watch_snapshot.json"

# ===== 검색 키워드 링크에 넣고 빌드 =====
def build_search_url_by_keyword(keyword: str, page: int = 1, page_num: int = 80) -> str:
    kw = quote(keyword.strip(), safe="")
    # keyword 파라미터만으로도 검색이 된다 하셨으니 가장 단순하게
    return f"{SEARCH_URL}?keyword={kw}&page={page}&pageNum={page_num}"

def ask_poll_seconds(default_minutes: int = 120) -> int:
    root = tk.Tk()
    root.withdraw()

    win = tk.Toplevel(root)
    style = ttk.Style()
    style.theme_use("clam")  # 기본보다 깔끔한 테마

    style.configure(
        "Primary.TButton",
        font=("맑은 고딕", 10, "bold"),
        padding=8,
        background="#2D6CDF",
        foreground="white"
    )

    style.map(
        "Primary.TButton",
        background=[("active", "#1E4FB5")]
    )

    style.configure(
        "Secondary.TButton",
        font=("맑은 고딕", 10),
        padding=8,
        background="#E5E5E5"
    )

    style.map(
        "Secondary.TButton",
        background=[("active", "#D0D0D0")]
    )
    win.title("감시 주기 설정")
    win.geometry("320x160")
    win.resizable(False, False)

    result = {"seconds": None}

    tk.Label(win, text="감시 주기(분 단위) 입력", font=("맑은 고딕", 11)).pack(pady=10)

    entry = tk.Entry(win, width=10)
    entry.pack()
    entry.insert(0, str(default_minutes))

    def ok():
        try:
            minutes = int(entry.get().strip())
            if minutes <= 0:
                raise ValueError
            result["seconds"] = minutes * 60
        except Exception:
            result["seconds"] = default_minutes * 60
        win.destroy()
        root.destroy()

    tk.Button(win, text="확인", width=10, command=ok).pack(pady=15)

    win.grab_set()
    win.wait_window()

    return result["seconds"] or (default_minutes * 60)

# ===== 감시 =====
def load_snapshot(path: str) -> dict:
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}

def save_snapshot(path: str, data: dict) -> None:
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)

def is_target_product(product_name: str, watch_keywords: list[str]) -> bool:
    pn = (product_name or "").strip()
    return any(k in pn for k in watch_keywords)

def build_option_key(opt: dict) -> str:
    # 1단 옵션은 optionSno가 신뢰도 높음
    sno = (opt.get("optionSno") or "").strip()
    if sno:
        return f"sno:{sno}"
    # 2단 옵션은 optionName으로 대체
    name = (opt.get("optionName") or "").strip()
    return f"name:{name}"

def normalize_option_state(opt: dict) -> dict:
    # 비교용 표준 형태
    return {
        "name": (opt.get("optionName") or "").strip(),
        "soldOut": (opt.get("isSoldOut") or "N"),
        "stock": (opt.get("optionStock") or "").strip(),
    }
def _to_int_stock(s: str) -> int | None:
    if s is None:
        return None
    s = str(s).strip()
    if s == "" or s.lower() in ("none", "null"):
        return None
    m = re.search(r"\d+", s)
    if not m:
        return None
    try:
        return int(m.group(0))
    except Exception:
        return None


def diff_product(prev: dict, cur: dict) -> tuple[list[str], list[str]]:
    """
    반환:
      - alerts: 알림(강조)용 메시지 (재고/품절 변화만)
      - infos: 정보용 메시지 (옵션추가/삭제 등)
    """
    alerts: list[str] = []
    infos: list[str] = []

    prev_opts = (prev or {}).get("options") or {}
    cur_opts = (cur or {}).get("options") or {}
    all_keys = set(prev_opts.keys()) | set(cur_opts.keys())

    for k in sorted(all_keys):
        a = prev_opts.get(k)
        b = cur_opts.get(k)

        # 옵션 추가/삭제는 "정보"로만
        if a is None and b is not None:
            infos.append(f"[옵션추가] {b['name']} (품절:{b['soldOut']}, 재고:{b['stock']})")
            continue
        if a is not None and b is None:
            infos.append(f"[옵션삭제] {a['name']}")
            continue

        # 둘 다 있으면 재고/품절 변화만 alerts로
        a_sold = (a.get("soldOut") or "N")
        b_sold = (b.get("soldOut") or "N")

        a_stock_raw = (a.get("stock") or "").strip()
        b_stock_raw = (b.get("stock") or "").strip()

        a_stock = _to_int_stock(a_stock_raw)
        b_stock = _to_int_stock(b_stock_raw)

        name = (b.get("name") or a.get("name") or "").strip()

        # 품절 전환/해제(강조)
        if a_sold != b_sold:
            if a_sold == "N" and b_sold == "Y":
                alerts.append(f"[품절전환] {name}")
                continue
            if a_sold == "Y" and b_sold == "N":
                if b_stock is not None:
                    alerts.append(f"[품절해제] {name} (재고 {b_stock})")
                else:
                    alerts.append(f"[품절해제] {name}")
                continue

        # 재고 증가/감소(강조) - 숫자로 비교 가능할 때만
        if a_stock is not None and b_stock is not None and a_stock != b_stock:
            if b_stock > a_stock:
                alerts.append(f"[재고증가] {name} {a_stock} → {b_stock}")
            else:
                alerts.append(f"[재고감소] {name} {a_stock} → {b_stock}")
            continue

        # 숫자 비교 불가한 케이스는 알림에서 제외(원하시면 infos로 이동 가능)
        # if a_stock_raw != b_stock_raw:
        #     infos.append(f"[재고변경] {name} '{a_stock_raw}' → '{b_stock_raw}'")

    return alerts, infos

def collect_goods_nos_for_keyword(sess: requests.Session, keyword: str, max_pages: int = 50, page_num: int = 80) -> list[str]:
    prev_signature = None
    repeat_count = 0

    all_nos: list[str] = []

    for page in range(1, max_pages + 1):
        url = build_search_url_by_keyword(keyword, page=page, page_num=page_num)
        r = sess.get(url, timeout=30)
        if r.status_code != 200:
            break
        r.encoding = "utf-8"

        goods_nos = extract_goods_nos_from_search(r.text)
        if not goods_nos:
            break

        signature = ",".join(goods_nos[:10])
        if signature == prev_signature:
            repeat_count += 1
        else:
            repeat_count = 0
        prev_signature = signature
        if repeat_count >= 2:
            break

        all_nos.extend(goods_nos)

        time.sleep(random.uniform(0.08, 0.18))

    # 중복 제거 + 순서 유지
    seen = set()
    uniq = []
    for n in all_nos:
        if n not in seen:
            seen.add(n)
            uniq.append(n)

    return uniq

# ===== 팝업 로그 ======
def start_log_window() -> tuple[tk.Tk, "queue.Queue[str]"]:
    q: "queue.Queue[str]" = queue.Queue()

    root = tk.Tk()
    root.title("Parker 재고 감시 로그")
    root.geometry("900x500")

    text = tk.Text(root, wrap="word")
    text.pack(fill="both", expand=True)

    def pump():
        try:
            while True:
                msg = q.get_nowait()
                text.insert("end", msg + "\n")
                text.see("end")
        except queue.Empty:
            pass
        root.after(200, pump)

    pump()
    return root, q

def log(q: "queue.Queue[str]", msg: str) -> None:
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    q.put(f"{ts}  {msg}")

# ===== 키워드 입력 받기 =====
def ask_watch_keywords(default: str = "트란지아,힐레베르그") -> list[str]:
    """
    감시할 상품명 키워드 입력 (쉼표/줄바꿈 구분)
    예: "트란지아, 힐레베르그"
    """
    root = tk.Tk()
    root.withdraw()

    win = tk.Toplevel(root)
    win.title("감시 키워드 입력")
    win.geometry("420x220")
    win.resizable(False, False)

    result = {"keywords": None}

    tk.Label(win, text="상품명에 포함될 키워드를 입력하세요", font=("맑은 고딕", 11)).pack(pady=10)
    tk.Label(win, text="쉼표(,) 또는 줄바꿈으로 여러 개 입력 가능", font=("맑은 고딕", 9)).pack()

    entry = tk.Text(win, height=5, width=45)
    entry.pack(pady=10)
    entry.insert("1.0", default)

    def ok():
        raw = entry.get("1.0", "end").strip()
        # 쉼표/줄바꿈 모두 지원
        parts = re.split(r"[,\n]+", raw)
        kws = [p.strip() for p in parts if p.strip()]
        result["keywords"] = kws
        win.destroy()
        root.destroy()

    def cancel():
        result["keywords"] = []
        win.destroy()
        root.destroy()

    btn_frame = tk.Frame(win)
    btn_frame.pack(pady=5)
    tk.Button(btn_frame, text="확인", width=10, command=ok).grid(row=0, column=0, padx=6)
    tk.Button(btn_frame, text="취소", width=10, command=cancel).grid(row=0, column=1, padx=6)

    win.protocol("WM_DELETE_WINDOW", cancel)
    win.grab_set()
    win.wait_window()

    return result["keywords"] or []


# ===== 감시 루프 =====
def monitor_loop(sess: requests.Session, q: "queue.Queue[str]", watch_keywords: list[str]):
    snapshot = load_snapshot(SNAPSHOT_PATH)
    log(q, f"감시 시작: 키워드={watch_keywords}, 주기={POLL_SECONDS}s, snapshot={SNAPSHOT_PATH}")

    first_cycle = True

    while True:
        try:
            changed_count = 0
            checked_products = 0

            # ✅ 1) 키워드별로 goodsNo 수집 (검색 기반)
            merged_goods_nos: list[str] = []
            seen = set()

            for kw in watch_keywords:
                nos = collect_goods_nos_for_keyword(sess, kw, max_pages=60, page_num=80)
                log(q, f"[SEARCH] '{kw}' 결과 goodsNo={len(nos)}")
                for n in nos:
                    if n not in seen:
                        seen.add(n)
                        merged_goods_nos.append(n)

            log(q, f"[TARGET] 전체(키워드 합집합) goodsNo={len(merged_goods_nos)}")

            # ✅ 2) 수집된 goodsNo만 상세 조회
            for goods_no in merged_goods_nos:
                view = sess.get(f"{VIEW_URL}?goodsNo={goods_no}", timeout=30)
                if view.status_code != 200:
                    continue

                view.encoding = "utf-8"
                data = parse_goods_view(sess, view.text, str(goods_no))

                checked_products += 1

                cur_opts = {}
                for opt in (data.get("options") or []):
                    key = build_option_key(opt)
                    cur_opts[key] = normalize_option_state(opt)

                cur_rec = {"productName": data.get("productName", ""), "options": cur_opts}
                prev_rec = snapshot.get(str(goods_no))

                alerts, infos = diff_product(prev_rec, cur_rec)

                if first_cycle:
                    snapshot[str(goods_no)] = cur_rec
                else:
                    if alerts:
                        changed_count += 1
                        log(q, f"=== ALERT goodsNo={goods_no} | {cur_rec['productName']} ===")
                        for m in alerts:
                            log(q, "  " + m)

                    # 필요 없으면 삭제 가능
                    if infos:
                        log(q, f"--- INFO goodsNo={goods_no} | {cur_rec['productName']} ---")
                        for m in infos:
                            log(q, "  " + m)

                    if alerts or infos:
                        snapshot[str(goods_no)] = cur_rec

                time.sleep(random.uniform(0.12, 0.28))

            if first_cycle:
                log(q, "[INIT] baseline captured (first cycle). next cycle will alert changes.")
                first_cycle = False

            log(q, f"[CYCLE DONE] 대상상품 체크={checked_products}, ALERT 변경상품={changed_count}")
            save_snapshot(SNAPSHOT_PATH, snapshot)

        except Exception as e:
            log(q, f"[ERROR] {type(e).__name__}: {e}")

        time.sleep(POLL_SECONDS)

# ===== 공통: goodsNo 추출 정규식 =====
_goodsno_re = re.compile(r"goods_view\.php\?goodsNo=(\d+)")


def wait_login_popup():
    root = tk.Tk()
    root.withdraw()
    messagebox.showinfo(
        title="로그인 필요",
        message=(
            "브라우저가 열렸습니다.\n\n"
            "1) 거래처 사이트에서 직접 로그인하세요.\n"
            "2) 상품 목록 화면이 정상적으로 보이면\n"
            "   [확인] 버튼을 눌러주세요."
        ),
    )
    root.destroy()


def show_login_failed_popup():
    root = tk.Tk()
    root.withdraw()
    messagebox.showwarning(
        title="로그인 확인 실패",
        message=(
            "로그인 상태가 확인되지 않았습니다.\n\n"
            "브라우저에서 로그인 상태(상품 목록 노출)를 확인한 뒤\n"
            "다시 [확인] 버튼을 눌러주세요."
        ),
    )
    root.destroy()

def ask_save_path(default_name="parkers_inventory.xlsx") -> str:
    root = tk.Tk()
    root.withdraw()

    path = filedialog.asksaveasfilename(
        title="엑셀 저장 위치 선택",
        defaultextension=".xlsx",
        initialfile=default_name,
        filetypes=[("Excel files", "*.xlsx")]
    )

    root.destroy()
    return path

def ask_open_path() -> str:
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="수정할 엑셀 파일 선택",
        filetypes=[("Excel files", "*.xlsx")]
    )
    root.destroy()
    return path

def ask_mode() -> str:
    root = tk.Tk()
    root.withdraw()

    win = tk.Toplevel(root)
    win.title("모드 선택")
    win.geometry("420x200")  # ✅ 가로 조금 늘림
    win.resizable(False, False)

    result = {"mode": None}

    tk.Label(
        win,
        text="실행 모드를 선택하세요",
        font=("맑은 고딕", 12),
        pady=15
    ).pack()

    def choose(m: str):
        result["mode"] = m
        win.destroy()
        root.destroy()

    btn_frame = tk.Frame(win)
    btn_frame.pack(pady=10)

    # ✅ 1줄에 2개 + 2줄 구성 추천
    ttk.Button(btn_frame, text="감시", width=18, command=lambda: choose("watch")).grid(row=0, column=0, padx=6, pady=4)
    ttk.Button(btn_frame, text="기존파일 업데이트", width=18, command=lambda: choose("update")).grid(row=0, column=1, padx=6, pady=4)

    ttk.Button(btn_frame, text="새 파일 전체수집", width=18, command=lambda: choose("new")).grid(row=1, column=0, padx=6, pady=4)
    ttk.Button(btn_frame, text="키워드 대상 수집", width=18, command=lambda: choose("keyword")).grid(row=1, column=1, padx=6, pady=4)

    win.protocol("WM_DELETE_WINDOW", lambda: choose("cancel"))
    win.grab_set()
    win.wait_window()

    return result["mode"] or "cancel"


def extract_goods_nos_from_search(html: str) -> list[str]:
    nos = _goodsno_re.findall(html)
    seen = set()
    out = []
    for n in nos:
        if n not in seen:
            seen.add(n)
            out.append(n)
    return out


def _extract_param(html: str, key: str) -> str | None:
    m = re.search(rf"'{re.escape(key)}'\s*:\s*'([^']*)'", html)
    if m:
        return m.group(1).strip()
    m = re.search(rf'"{re.escape(key)}"\s*:\s*"([^"]*)"', html)
    if m:
        return m.group(1).strip()
    return None



def _post_goods_ps(sess: requests.Session, payload: dict, option_vals: list[str]) -> dict:
    """
    goods_ps.php에 option_select 요청을 보내고 JSON dict를 반환합니다.
    optionVal[]를 리스트로 안전하게 전달하기 위해 tuples를 사용합니다.
    """
    data = []
    for k, v in payload.items():
        data.append((k, v))
    for v in option_vals:
        data.append(("optionVal[]", v))

    r = sess.post(GOODS_PS_URL, data=data, timeout=30)
    if r.status_code != 200:
        return {}

    # JSON 파싱 (content-type이 text/html이어도 JSON 문자열일 수 있음)
    try:
        return r.json()
    except Exception:
        txt = (r.text or "").strip()
        if txt.startswith("{") and txt.endswith("}"):
            try:
                return json.loads(txt)
            except Exception:
                return {}
        return {}

def goods_ps_option_select(sess: requests.Session, goods_no: str, option_key: int, option_vals: list[str], mileage_fl: str = "c") -> dict:
    """
    mode=option_select 호출
    - option_key: 0,1,2...
    - option_vals: 선택한 옵션 값들 (순서대로)
    """
    payload = {
        "mode": "option_select",
        "optionKey": str(option_key),
        "goodsNo": str(goods_no),
        "mileageFl": mileage_fl,   # 캡처: ' c' 로 보이지만 실제 전송은 'c'인 경우가 많음
    }
    return _post_goods_ps(sess, payload, option_vals)

def extract_option_levels_from_view(html: str) -> list[list[str]]:
    """
    goods_view.php HTML에서 optionNo_0, optionNo_1 ... 의 옵션 값을 추출합니다.
    반환: [[1단값들], [2단값들], ...]
    - 2단값들은 보통 초기엔 비활성/미표시일 수 있으므로, 전수 수집은 goods_ps 응답 nextOption로 진행합니다.
    """
    soup = BeautifulSoup(html, "html.parser")

    levels = []
    i = 0
    while True:
        sel = soup.select_one(f'select[name="optionNo_{i}"]')
        if not sel:
            break
        vals = []
        for opt in sel.select("option"):
            v = (opt.get("value") or "").strip()
            if not v:
                continue
            # placeholder (= 색상 선택 =) 형태 제거
            if "선택" in opt.get_text(" ", strip=True):
                continue
            # 2단 옵션의 option 태그 value는 보통 "XS" 같은 순수 값
            vals.append(v)
        # 중복 제거 유지
        uniq = []
        seen = set()
        for v in vals:
            if v not in seen:
                seen.add(v)
                uniq.append(v)
        levels.append(uniq)
        i += 1

    return levels

def _pick_stock_from_goods_ps(j: dict) -> str:
    """
    1단 옵션에서 goods_ps 응답으로부터 재고를 최대한 안전하게 뽑습니다.
    stockCnt 형태가 list/str/int 등으로 올 수 있어 방어적으로 처리합니다.
    """
    sc = j.get("stockCnt")

    # list 형태: ["3"] 또는 [3] 또는 ["0", ...]
    if isinstance(sc, list) and len(sc) > 0:
        return str(sc[0]).strip()

    # 단일 값 형태
    if sc is None:
        return ""
    return str(sc).strip()

def extract_1depth_options_from_optionSnoInput(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    sel = soup.select_one('select[name="optionSnoInput"]')
    if not sel:
        return []

    out: list[dict] = []
    for opt in sel.select("option"):
        v = (opt.get("value") or "").strip()
        text = opt.get_text(" ", strip=True)

        # placeholder 제거 (괄호로 우선순위 고정)
        if (not v) or (("옵션" in text) and ("가격" in text)):
            continue

        option_sno = ""
        option_price = ""
        stock = ""
        option_name = ""

        # 1) option value 파싱: "6582||6500||||14^|^BLACK BALLISTIC"
        parts = v.split("||")
        if len(parts) >= 2:
            option_sno = parts[0].strip()
            option_price = parts[1].strip()

        if "||||" in v:
            tail = v.split("||||", 1)[1].strip()
            if "^|^" in tail:
                stock_part, name_part = tail.split("^|^", 1)
                stock = stock_part.strip()
                option_name = name_part.strip()
            else:
                option_name = tail.strip()

        # 이름 보조 추출
        if not option_name:
            option_name = text.split(":")[0].strip()

        # 2) ✅ 품절 판정: disabled 속성 또는 텍스트에 '품절' 포함 시 무조건 품절
        is_disabled = opt.has_attr("disabled")
        text_soldout = ("품절" in text)

        soldout = is_disabled or text_soldout

        out.append({
            "optionName": option_name,
            "optionSno": option_sno,
            # 품절이면 재고는 비움(숫자가 있어도 구매불가이므로)
            "optionStock": "" if soldout else stock,
            "isSoldOut": "Y" if soldout else "N",
            "optionPrice": option_price,
        })

    return out

def parse_goods_view(sess: requests.Session, html: str, goods_no: str) -> dict:
    soup = BeautifulSoup(html, "html.parser")

    name_el = soup.select_one("div.goods-header h2")
    product_name = (name_el.get_text(" ", strip=True) if name_el else "").strip()

    option_fl = _extract_param(html, "setOptionFl")  # 'y' or 'n'
    stock_cnt = _extract_param(html, "setStockCnt")
    goods_price = _extract_param(html, "setGoodsPrice")

    # fixedPrice(소비자가) 추출
    fixed_price = ""
    for li in soup.select(".item ul li"):
        strong = li.select_one("strong")
        if not strong:
            continue
        label = strong.get_text(" ", strip=True)
        if "소비자가" in label:
            sp = li.select_one("del span")
            if sp:
                fixed_price = sp.get_text(" ", strip=True).replace(",", "").strip()
            break

    # 옵션상품 아니면 종료
    if (option_fl or "").strip().lower() != "y":
        return {
            "goodsNo": goods_no,
            "productName": product_name,
            "optionFl": (option_fl or ""),
            "stockCnt": (stock_cnt or ""),
            "goodsPrice": (goods_price or ""),
            "fixedPrice": fixed_price,
            "options": [],
        }

    # HTML에서 optionNo_0, optionNo_1 ... 값 목록 확인
    levels = extract_option_levels_from_view(html)

    # ✅ 1단 옵션(옵션 select가 optionSnoInput 1개) 케이스 우선 처리
    one_depth_opts = extract_1depth_options_from_optionSnoInput(html)
    if one_depth_opts:
        return {
            "goodsNo": goods_no,
            "productName": product_name,
            "optionFl": (option_fl or ""),
            "stockCnt": (stock_cnt or ""),
            "goodsPrice": (goods_price or ""),
            "fixedPrice": fixed_price,
            "options": [
                # optionSummary 생성 로직이 optionName/optionStock/isSoldOut만 쓰니까 그 3개는 유지
                {"optionName": o["optionName"], "optionSno": o["optionSno"], "optionStock": o["optionStock"], "isSoldOut": o["isSoldOut"]}
                for o in one_depth_opts
            ],
        }
    
    # optionFl=y인데 select가 없으면(희귀) 옵션 수집 불가로 반환
    if not levels or not levels[0]:
        return {
            "goodsNo": goods_no,
            "productName": product_name,
            "optionFl": (option_fl or ""),
            "stockCnt": (stock_cnt or ""),
            "goodsPrice": (goods_price or ""),
            "fixedPrice": fixed_price,
            "options": [],
        }

    options: list[dict] = []

    # -------------------------
    # A) 1단 옵션만 있는 경우 (수정 버전)
    # -------------------------
    if len(levels) == 1:
        vlist = levels[0]
        options: list[dict] = []

        for v in vlist:
            j1 = goods_ps_option_select(sess, goods_no, option_key=0, option_vals=[v], mileage_fl="c")

            stock = _pick_stock_from_goods_ps(j1)
            # optionSno가 내려오는 상품이면 같이 기록
            option_sno = (j1.get("optionSno") or "").strip()

            # "0"이면 품절 처리, 재고가 아예 비면 미확인 처리
            if stock == "0":
                options.append({
                    "optionName": str(v).strip(),
                    "optionSno": option_sno,
                    "optionStock": "",
                    "isSoldOut": "Y",
                })
            else:
                options.append({
                    "optionName": str(v).strip(),
                    "optionSno": option_sno,
                    "optionStock": stock,  # ""이면 미확인 상태로 남음
                    "isSoldOut": "N" if stock != "" else "N",
                })

            time.sleep(random.uniform(0.1, 0.25))  # 1단 옵션 반복 호출이므로 약간 더 짧게

        return {
            "goodsNo": goods_no,
            "productName": product_name,
            "optionFl": (option_fl or ""),
            "stockCnt": (stock_cnt or ""),
            "goodsPrice": (goods_price or ""),
            "fixedPrice": fixed_price,
            "options": options,
        }
    # -------------------------
    # B) 2단 이상 옵션인 경우
    # -------------------------
    for v1 in levels[0]:
        j1 = goods_ps_option_select(sess, goods_no, option_key=0, option_vals=[v1], mileage_fl="c")
        next_opts = j1.get("nextOption") or []
        stock_cnts = j1.get("stockCnt") or []

        pairs = []
        if isinstance(next_opts, list) and isinstance(stock_cnts, list) and len(next_opts) == len(stock_cnts):
            pairs = list(zip(next_opts, stock_cnts))
        elif isinstance(next_opts, list):
            pairs = [(x, "") for x in next_opts]

        for v2, st in pairs:
            stock = str(st).strip() if st is not None else ""
            soldout = (stock == "0" and stock != "")

            option_sno = ""
            if not soldout:
                j2 = goods_ps_option_select(sess, goods_no, option_key=1, option_vals=[v1, v2], mileage_fl="c")
                option_sno = (j2.get("optionSno") or "").strip()

                sc2 = j2.get("stockCnt")
                if isinstance(sc2, list) and sc2:
                    stock2 = str(sc2[0]).strip()
                    soldout = (stock2 == "0")
                    stock = stock2

            options.append({
                "optionName": f"{v1}/{v2}",
                "optionSno": option_sno,
                "optionStock": "" if soldout else stock,
                "isSoldOut": "Y" if soldout else "N",
            })

    return {
        "goodsNo": goods_no,
        "productName": product_name,
        "optionFl": (option_fl or ""),
        "stockCnt": (stock_cnt or ""),
        "goodsPrice": (goods_price or ""),
        "fixedPrice": fixed_price,
        "options": options,
    }

def assert_logged_in(sess: requests.Session) -> bool:
    test_url = f"{SEARCH_URL}?keyword=0&page=1&pageNum=80"
    r = sess.get(test_url, timeout=30)
    if r.status_code != 200:
        return False
    r.encoding = "utf-8"
    return bool(_goodsno_re.search(r.text))


def interactive_login_and_get_session() -> requests.Session:
    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()))

    # 처음부터 검색 페이지를 열어두면, 로그인 필요 시 자동으로 로그인 페이지로 이동합니다.
    driver.get(f"{SEARCH_URL}?keyword=0&page=1&pageNum=80")

    while True:
        wait_login_popup()

        sess = requests.Session()
        for c in driver.get_cookies():
            sess.cookies.set(c["name"], c["value"], domain=COOKIE_DOMAIN)

        try:
            if assert_logged_in(sess):
                break
        except requests.RequestException:
            pass

        show_login_failed_popup()

    driver.quit()
    return sess


def open_or_create_workbook(out_path: str):
    if os.path.exists(out_path):
        wb = load_workbook(out_path)
        if "Products" in wb.sheetnames:
            ws_p = wb["Products"]
        else:
            ws_p = wb.active
            ws_p.title = "Products"

        expected_new = ["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "fixedPrice", "optionSummary", "fetchedAt"]
        expected_old = ["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "optionSummary", "fetchedAt"]

        header = [cell.value for cell in ws_p[1]] if ws_p.max_row >= 1 else []

        if header == expected_new:
            pass
        elif header == expected_old:
            insert_at = 7  # goodsPrice 다음에 fixedPrice
            ws_p.insert_cols(insert_at)
            ws_p.cell(row=1, column=insert_at).value = "fixedPrice"
            for col, name in enumerate(expected_new, start=1):
                ws_p.cell(row=1, column=col).value = name
            wb.save(out_path)
            print("[INFO] upgraded existing file: added fixedPrice column")
        else:
            raise RuntimeError(
                f"기존 파일 헤더가 예상과 다릅니다.\n기존: {header}\n예상(신규): {expected_new}\n예상(구버전): {expected_old}"
            )

        existing = set()
        for row in ws_p.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            goods_no = row[1]
            if goods_no is None:
                continue
            existing.add(str(goods_no).strip())

        print(f"[INFO] append mode: existing goodsNo={len(existing)}")
        return wb, ws_p, existing 

    # 새 파일 생성
    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "Products"
    ws_p.append(["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "fixedPrice", "optionSummary", "fetchedAt"])
    print("[INFO] new file mode")
    return wb, ws_p, set()

# 상품 전수 조사
def crawl_to_excel(sess: requests.Session, out_path: str, max_pages: int = 300):
    wb, ws_p, existing_goods = open_or_create_workbook(out_path)

    vendor = "parkers"
    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    prev_signature = None
    repeat_count = 0

    try:
        for page in range(1, max_pages + 1):
            url = f"{SEARCH_URL}?keyword=0&page={page}&pageNum=80"
            try:
                r = sess.get(url, timeout=30)
            except requests.RequestException as e:
                print(f"[STOP] search page={page} request error: {e}")
                break

            if r.status_code != 200:
                print(f"[STOP] search page={page} status={r.status_code}")
                break

            r.encoding = "utf-8"
            goods_nos = extract_goods_nos_from_search(r.text)

            if not goods_nos:
                print(f"[DONE] no goods on page={page}")
                break
            signature = ",".join(goods_nos[:10])

            if signature == prev_signature:
                repeat_count += 1
            else:
                repeat_count = 0
            prev_signature = signature

            if repeat_count >= 2:
                print(f"[DONE] repeated page detected at page={page}. stop.")
                wb.save(out_path)
                return
            print(f"[LIST] page={page} goods={len(goods_nos)}")

            for idx, goods_no in enumerate(goods_nos, start=1):
                if str(goods_no) in existing_goods:
                    continue
                try:
                    view = sess.get(f"{VIEW_URL}?goodsNo={goods_no}", timeout=30)
                except requests.RequestException as e:
                    print(f"[SKIP] goodsNo={goods_no} request error: {e}")
                    continue

                if view.status_code != 200:
                    print(f"[SKIP] goodsNo={goods_no} status={view.status_code}")
                    continue

                view.encoding = "utf-8"
                data = parse_goods_view(sess, view.text, goods_no)

                option_summary_lines = []

                if data["optionFl"] == "y":
                    for opt in data["options"]:
                        name = (opt.get("optionName") or "").strip()
                        sold = (opt.get("isSoldOut") == "Y")
                        stock = (opt.get("optionStock") or "").strip()

                        if sold:
                            option_summary_lines.append(f"{name} [품절]")
                        else:
                            # 재고가 비면 미확인
                            if stock != "":
                                option_summary_lines.append(f"{name} - 재고({stock})")
                            else:
                                option_summary_lines.append(f"{name} - 재고(미확인)")

                option_summary = "\n".join(option_summary_lines)
                
                ws_p.append(
                    [
                        vendor,
                        data["goodsNo"],
                        data["productName"],
                        data["optionFl"],
                        data["stockCnt"],
                        data["goodsPrice"],
                        data.get("fixedPrice", ""),
                        option_summary,
                        fetched_at,
                    ]
                )
                existing_goods.add(str(goods_no))

                # 차단/부하 방지: 상세 요청마다 짧은 딜레이
                time.sleep(random.uniform(0.2, 0.6))

            if page % 2 == 0:
                wb.save(out_path)
                print(f"[SAVE] page={page} done (every 2 pages) -> {out_path}")

    except KeyboardInterrupt:
        print("[INTERRUPT] saving partial file...")
        wb.save(out_path)
        print(f"[OK] partial saved: {out_path}")
        return
    
    wb.save(out_path)
    print(f"[OK] saved: {out_path}")


# 기존 파일 업데이트
def fill_missing_option_summaries(sess: requests.Session, xlsx_path: str, save_as: str | None = None):
    """
    기존 엑셀을 열어서 optionSummary가 비어있는(옵션 1단/옵션상품 포함) 행만 다시 조회 후 채워 넣습니다.
    save_as가 있으면 새 파일로 저장, 없으면 원본 덮어쓰기 저장.
    """
    wb = load_workbook(xlsx_path)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.active

    # 헤더에서 컬럼 인덱스 찾기 (1-based)
    header = [c.value for c in ws[1]]
    def col(name: str) -> int:
        if name not in header:
            raise RuntimeError(f"엑셀 헤더에 '{name}' 컬럼이 없습니다. 현재 헤더: {header}")
        return header.index(name) + 1

    c_goodsNo = col("goodsNo")
    c_optionFl = col("optionFl")
    c_optionSummary = col("optionSummary")

    # fixedPrice 컬럼이 있으면 업데이트 시 함께 채우고 싶을 수도 있으니(선택)
    c_fixedPrice = header.index("fixedPrice") + 1 if "fixedPrice" in header else None

    updated = 0
    checked = 0

    for r in range(2, ws.max_row + 1):
        goods_no = ws.cell(r, c_goodsNo).value
        option_fl = str(ws.cell(r, c_optionFl).value or "").strip().lower()
        option_summary = str(ws.cell(r, c_optionSummary).value or "").strip()

        if option_fl != "y":
            continue
        if option_summary != "":
            continue

        checked += 1
        goods_no_str = str(goods_no).strip()

        try:
            view = sess.get(f"{VIEW_URL}?goodsNo={goods_no_str}", timeout=30)
            if view.status_code != 200:
                print(f"[SKIP] goodsNo={goods_no_str} view status={view.status_code}")
                continue
            view.encoding = "utf-8"
            data = parse_goods_view(sess, view.text, goods_no_str)
        except requests.RequestException as e:
            print(f"[SKIP] goodsNo={goods_no_str} request error: {e}")
            continue

        # optionSummary 생성
        lines = []
        if (data.get("optionFl") or "") == "y":
            for opt in data.get("options") or []:
                name = (opt.get("optionName") or "").strip()
                sold = (opt.get("isSoldOut") == "Y")
                stock = (opt.get("optionStock") or "").strip()

                if not name:
                    continue

                if sold:
                    lines.append(f"{name} [품절]")
                else:
                    if stock != "":
                        lines.append(f"{name} - 재고({stock})")
                    else:
                        lines.append(f"{name} - 재고(미확인)")

        new_summary = "\n".join(lines).strip()

        # 채우기
        if new_summary:
            ws.cell(r, c_optionSummary).value = new_summary
            updated += 1

        # fixedPrice도 비어있고 parse에서 뽑혔으면 채우기(선택)
        if c_fixedPrice is not None:
            cell_val = ws.cell(r, c_fixedPrice).value
            cur_fp = str(cell_val).strip() if cell_val is not None else ""
            fp = str(data.get("fixedPrice") or "").strip()

            if cur_fp in ("", "0", "0.0") and fp:
                ws.cell(r, c_fixedPrice).value = fp

        if checked % 20 == 0:
            out = save_as or xlsx_path
            wb.save(out)
            print(f"[SAVE] checked={checked}, updated={updated} -> {out}")

        time.sleep(random.uniform(0.2, 0.6))

    out = save_as or xlsx_path
    wb.save(out)
    print(f"[OK] fill done. checked={checked}, updated={updated} -> {out}")

# 키워드 기반 정보 수집
def crawl_keyword_to_excel(
    sess: requests.Session,
    out_path: str,
    watch_keywords: list[str],
    max_pages_per_keyword: int = 60,
    page_num: int = 80
):
    """
    ✅ 키워드 대상 수집 모드
    - 감시 모드의 검색(goodsNo 수집) 로직 사용
    - 전체수집과 동일하게 parse_goods_view로 상세 수집 후 엑셀 저장
    """
    wb, ws_p, existing_goods = open_or_create_workbook(out_path)

    vendor = "parkers"
    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # 1) 키워드별 goodsNo 합집합 수집 (감시 로직)
    merged_goods_nos: list[str] = []
    seen = set()

    for kw in watch_keywords:
        nos = collect_goods_nos_for_keyword(sess, kw, max_pages=max_pages_per_keyword, page_num=page_num)
        print(f"[SEARCH] '{kw}' 결과 goodsNo={len(nos)}")
        for n in nos:
            if n not in seen:
                seen.add(n)
                merged_goods_nos.append(n)

    print(f"[TARGET] 전체(키워드 합집합) goodsNo={len(merged_goods_nos)}")

    # 2) 수집된 goodsNo만 상세 조회 후 엑셀 반영
    saved = 0
    skipped_exist = 0

    for i, goods_no in enumerate(merged_goods_nos, start=1):
        if str(goods_no) in existing_goods:
            skipped_exist += 1
            continue

        try:
            view = sess.get(f"{VIEW_URL}?goodsNo={goods_no}", timeout=30)
        except requests.RequestException as e:
            print(f"[SKIP] goodsNo={goods_no} request error: {e}")
            continue

        if view.status_code != 200:
            print(f"[SKIP] goodsNo={goods_no} status={view.status_code}")
            continue

        view.encoding = "utf-8"
        data = parse_goods_view(sess, view.text, str(goods_no))

        option_summary_lines = []
        if (data.get("optionFl") or "").strip().lower() == "y":
            for opt in data.get("options") or []:
                name = (opt.get("optionName") or "").strip()
                sold = (opt.get("isSoldOut") == "Y")
                stock = (opt.get("optionStock") or "").strip()

                if not name:
                    continue

                if sold:
                    option_summary_lines.append(f"{name} [품절]")
                else:
                    if stock != "":
                        option_summary_lines.append(f"{name} - 재고({stock})")
                    else:
                        option_summary_lines.append(f"{name} - 재고(미확인)")

        option_summary = "\n".join(option_summary_lines)

        ws_p.append(
            [
                vendor,
                data.get("goodsNo", str(goods_no)),
                data.get("productName", ""),
                data.get("optionFl", ""),
                data.get("stockCnt", ""),
                data.get("goodsPrice", ""),
                data.get("fixedPrice", ""),
                option_summary,
                fetched_at,
            ]
        )
        existing_goods.add(str(goods_no))
        saved += 1

        # 저장/딜레이
        if i % 30 == 0:
            wb.save(out_path)
            print(f"[SAVE] i={i}, saved={saved}, skipped_exist={skipped_exist} -> {out_path}")

        time.sleep(random.uniform(0.2, 0.6))

    wb.save(out_path)
    print(f"[OK] keyword crawl saved={saved}, skipped_exist={skipped_exist} -> {out_path}")



if __name__ == "__main__":
    mode = ask_mode()
    if mode == "cancel":
        print("[CANCEL] mode select canceled by user")
        exit(0)

    sess = interactive_login_and_get_session()

    if mode == "watch":
        keywords = ask_watch_keywords(default="트란지아,힐레베르그")
        if not keywords:
            exit(0)

        poll_seconds = ask_poll_seconds(default_minutes=120)

        global POLL_SECONDS
        POLL_SECONDS = poll_seconds

        root, q = start_log_window()
        t = threading.Thread(target=monitor_loop, args=(sess, q, keywords), daemon=True)
        t.start()
        root.mainloop()
        exit(0)


    if mode == "update":
        in_path = ask_open_path()
        if not in_path:
            print("[CANCEL] open canceled by user")
            exit(0)

        out_path = ask_save_path(default_name="parkers_inventory_filled.xlsx")
        if not out_path:
            print("[CANCEL] save canceled by user")
            exit(0)

        print("[INFO] update input :", in_path)
        print("[INFO] update output:", out_path)

        fill_missing_option_summaries(sess, xlsx_path=in_path, save_as=out_path)
        exit(0)

    if mode == "new":
        out_path = ask_save_path(default_name="parkers_inventory.xlsx")
        if not out_path:
            print("[CANCEL] save canceled by user")
            exit(0)

        print("[INFO] new crawl output:", out_path)

        crawl_to_excel(sess, out_path=out_path, max_pages=300)
        exit(0)

    if mode == "keyword":
        keywords = ask_watch_keywords(default="트란지아,힐레베르그")
        if not keywords:
            print("[CANCEL] keyword empty")
            exit(0)

        out_path = ask_save_path(default_name="parkers_inventory_keyword.xlsx")
        if not out_path:
            print("[CANCEL] save canceled by user")
            exit(0)

        print("[INFO] keyword crawl output:", out_path)
        print("[INFO] keywords:", keywords)

        crawl_keyword_to_excel(
            sess,
            out_path=out_path,
            watch_keywords=keywords,
            max_pages_per_keyword=60,
            page_num=80
        )
        exit(0)

    print(f"[CANCEL] unknown mode: {mode}")
    exit(0)
