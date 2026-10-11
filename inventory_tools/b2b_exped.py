import json
import math
import os
import queue
import random
import re
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urljoin, urlparse, urlunparse
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.common.exceptions import NoAlertPresentException, TimeoutException
from selenium.webdriver.support.ui import WebDriverWait

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

BASE_URL = "https://www.exped.co.kr"
CALCULATOR_URL = BASE_URL + "/exec/front/shop/CalculatorProduct"
REQUEST_QUANTITY = 999
SAVE_EVERY_PRODUCTS = 5
MAX_EMPTY_PAGE_COUNT = 2
SLEEP_BETWEEN_PRODUCTS = (0.15, 0.35)
SLEEP_BETWEEN_STOCK_REQUESTS = (0.08, 0.18)

CATEGORY_URLS = [
    "https://www.exped.co.kr/product/list.html?cate_no=24",
    "https://www.exped.co.kr/product/list.html?cate_no=25",
    "https://www.exped.co.kr/product/list.html?cate_no=26",
    "https://www.exped.co.kr/product/list.html?cate_no=27",
    "https://www.exped.co.kr/product/list.html?cate_no=42",
    "https://www.exped.co.kr/product/list.html?cate_no=28",
    "https://www.exped.co.kr/product/list.html?cate_no=43",
    "https://www.exped.co.kr/product/list.html?cate_no=44",
]

HEADERS = ["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "fixedPrice", "optionSummary", "fetchedAt", "debugUrls"]
VENDOR = "exped"
LOG_QUEUE = None
STOCK_DRIVER = None

try:
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
except Exception:
    pass


def log(message: str) -> None:
    print(message, flush=True)
    if LOG_QUEUE is not None:
        try:
            LOG_QUEUE.put(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  {message}")
        except Exception:
            pass


def start_log_window() -> tuple[tk.Tk, "queue.Queue[str]"]:
    q: "queue.Queue[str]" = queue.Queue()
    root = tk.Tk()
    root.title("EXPED 재고 수집 로그")
    root.geometry("980x560")

    text = tk.Text(root, wrap="word")
    text.pack(fill="both", expand=True)

    def pump():
        try:
            while True:
                msg = q.get_nowait()
                text.insert("end", msg + "\n")
                text.see("end")
        except queue.Empty:
            pass
        root.after(200, pump)

    pump()
    return root, q


def ask_mode() -> str:
    root = tk.Tk()
    root.withdraw()

    win = tk.Toplevel(root)
    win.title("모드 선택")
    win.geometry("420x190")
    win.resizable(False, False)

    result = {"mode": "cancel"}
    tk.Label(win, text="실행 모드를 선택하세요", font=("맑은 고딕", 12), pady=15).pack()

    def choose(mode: str):
        result["mode"] = mode
        win.destroy()
        root.destroy()

    btn_frame = tk.Frame(win)
    btn_frame.pack(pady=10)
    ttk.Button(btn_frame, text="새 파일 전체수집", width=18, command=lambda: choose("new")).grid(row=0, column=0, padx=6, pady=4)
    ttk.Button(btn_frame, text="기존파일 업데이트", width=18, command=lambda: choose("update")).grid(row=0, column=1, padx=6, pady=4)
    ttk.Button(btn_frame, text="취소", width=18, command=lambda: choose("cancel")).grid(row=1, column=0, columnspan=2, padx=6, pady=4)

    win.protocol("WM_DELETE_WINDOW", lambda: choose("cancel"))
    win.grab_set()
    win.wait_window()
    return result["mode"]


def ask_save_path(default_name: str = "exped_inventory.xlsx") -> str:
    root = tk.Tk()
    root.withdraw()
    path = filedialog.asksaveasfilename(
        title="엑셀 저장 위치 선택",
        defaultextension=".xlsx",
        initialfile=default_name,
        filetypes=[("Excel files", "*.xlsx")],
    )
    root.destroy()
    return path


def ask_open_path() -> str:
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(title="수정할 엑셀 파일 선택", filetypes=[("Excel files", "*.xlsx")])
    root.destroy()
    return path


def show_popup(title: str, message: str, error: bool = False) -> None:
    root = tk.Tk()
    root.withdraw()
    if error:
        messagebox.showerror(title, message)
    else:
        messagebox.showinfo(title, message)
    root.destroy()


def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def digits_only(value: str | None) -> str:
    return re.sub(r"[^0-9]", "", value or "")


def absolute_url(href: str | None) -> str:
    return urljoin(BASE_URL, href or "").split("#")[0]


def extract_product_no(url: str) -> str:
    patterns = [r"product_no=(\d+)", r"/product/[^/]+/(\d+)/", r"anchorBoxId_(\d+)"]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return ""


def extract_cate_no(url: str) -> str:
    patterns = [r"cate_no=(\d+)", r"/category/(?:[^/]+/)?(\d+)/", r"/category/[^/]+/(\d+)/"]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return ""


def make_page_url(category_url: str, page: int) -> str:
    parsed = urlparse(category_url)
    query = parse_qs(parsed.query)
    query["page"] = [str(page)]
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, urlencode(query, doseq=True), ""))


def product_no_to_code(product_no: str) -> str:
    number = int(product_no)
    letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    if number < 26 * 26:
        return f"P00000{letters[number // 26]}{letters[number % 26]}"
    n = number
    chunks = []
    while n >= 0:
        chunks.append(letters[n % 26])
        n = n // 26 - 1
        if n < 0:
            break
    suffix = "".join(reversed(chunks))
    return "P" + suffix.rjust(7, "0")


def option_suffix(index: int) -> str:
    letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    if index < 26:
        return f"000{letters[index]}"
    first = letters[(index // 26) - 1]
    second = letters[index % 26]
    return f"00{first}{second}"


def option_suffix_candidates(limit: int = 52) -> list[str]:
    return [option_suffix(i) for i in range(limit)]


def is_useful_stock_response(stock_info: dict) -> bool:
    result = stock_info.get("result")
    stock = clean_text(stock_info.get("stock"))
    return result in ("stock_number", "soldout_message", "quantity_alert_limit", "quantity_alert_stock_limit") and stock != ""


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
        "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
        "Connection": "keep-alive",
    })
    return session


def get_text(session: requests.Session, url: str, referer: str | None = None) -> str:
    headers = {"Referer": referer or BASE_URL}
    response = session.get(url, headers=headers, timeout=30)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or "utf-8"
    return response.text


def parse_last_page(html: str) -> int:
    soup = BeautifulSoup(html, "html.parser")
    pages = [1]
    for a in soup.select(".xans-product-normalpaging a[href*='page=']"):
        query = parse_qs(urlparse(a.get("href", "")).query)
        page_values = query.get("page") or []
        if page_values and page_values[0].isdigit():
            pages.append(int(page_values[0]))
    count_el = soup.select_one(".prdCount strong")
    count_text = digits_only(count_el.get_text(" ", strip=True)) if count_el else ""
    if count_text:
        pages.append(max(1, math.ceil(int(count_text) / 12)))
    return max(pages)


def parse_product_links(html: str, category_url: str, list_page_url: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    root = soup.select_one("#contents") or soup
    products = []
    seen = set()

    for li in root.select("ul.prdList > li[id^='anchorBoxId_']"):
        a = li.select_one(".prdImg a[href*='/product/']") or li.select_one("strong.name a[href*='/product/']")
        if not a:
            continue

        product_url = absolute_url(a.get("href"))
        product_no = extract_product_no(product_url) or extract_product_no(li.get("id", ""))
        cate_no = extract_cate_no(product_url) or extract_cate_no(category_url)
        if not product_no or product_no in seen:
            continue

        name_el = li.select_one("strong.name a span:not(.title) span") or li.select_one("strong.name a")
        list_name = clean_text(name_el.get_text(" ", strip=True)) if name_el else ""
        if not list_name:
            img = a.select_one("img")
            list_name = clean_text(img.get("alt")) if img else ""

        price_el = li.select_one(".description[ec-data-price]")
        list_price = price_el.get("ec-data-price", "") if price_el else ""
        if not list_price:
            price_text_el = li.select_one("li[class*='판매가'] span:not(.title)")
            list_price = digits_only(price_text_el.get_text(" ", strip=True)) if price_text_el else ""

        list_soldout = "Y" if li.select_one(".promotion img[alt*='품절']") else "N"
        products.append({
            "category_url": category_url,
            "list_page_url": list_page_url,
            "product_url": product_url,
            "product_no": product_no,
            "cate_no": cate_no,
            "list_name": list_name,
            "list_price": list_price,
            "list_soldout": list_soldout,
        })
        seen.add(product_no)

    return products


def iter_category_products(session: requests.Session):
    emitted = set()
    for category_index, category_url in enumerate(CATEGORY_URLS, start=1):
        log(f"[CATEGORY {category_index}/{len(CATEGORY_URLS)}] {category_url}")
        try:
            first_html = get_text(session, category_url)
        except Exception as exc:
            log(f"  [SKIP] 카테고리 접근 실패: {type(exc).__name__}: {exc}")
            continue

        last_page = parse_last_page(first_html)
        empty_count = 0

        for page in range(1, last_page + 1):
            page_url = make_page_url(category_url, page)
            try:
                html = first_html if page == 1 else get_text(session, page_url, referer=category_url)
            except Exception as exc:
                log(f"  [SKIP] page={page} 접근 실패: {type(exc).__name__}: {exc}")
                continue

            links = parse_product_links(html, category_url, page_url)
            log(f"  [LIST] page={page}/{last_page}, 상품={len(links)}")

            if not links:
                empty_count += 1
                if empty_count >= MAX_EMPTY_PAGE_COUNT:
                    log("  [STOP] 빈 페이지가 반복되어 해당 카테고리를 종료합니다.")
                    break
            else:
                empty_count = 0

            for link in links:
                product_no = link["product_no"]
                if product_no in emitted:
                    continue
                emitted.add(product_no)
                yield link

            time.sleep(random.uniform(0.08, 0.18))


def parse_product_meta(html: str, product_link: dict) -> dict:
    soup = BeautifulSoup(html, "html.parser")

    name_el = soup.select_one(".infoArea .headingArea h1") or soup.select_one("h1") or soup.select_one("tr.상품명 td")
    product_name = clean_text(name_el.get_text(" ", strip=True)) if name_el else product_link.get("list_name", "")

    price_el = soup.select_one("#span_product_price_text") or soup.select_one(".price-original")
    goods_price = digits_only(price_el.get_text(" ", strip=True)) if price_el else digits_only(product_link.get("list_price", ""))

    fixed_price = ""
    custom_el = soup.select_one("#ec-product-price-info[ec-data-custom]")
    if custom_el:
        custom_value = digits_only(custom_el.get("ec-data-custom", ""))
        if custom_value and custom_value != "0":
            fixed_price = custom_value

    product_code = ""
    image_zoom = soup.select_one("a[href*='image_zoom2'][alt^='P']")
    if image_zoom:
        product_code = clean_text(image_zoom.get("alt"))
    if not product_code:
        match = re.search(r"['\"](P[0-9A-Z]{7})['\"]", html)
        product_code = match.group(1) if match else ""
    if not product_code and product_link.get("product_no", "").isdigit():
        product_code = product_no_to_code(product_link["product_no"])

    return {
        "product_name": product_name,
        "goods_price": goods_price,
        "fixed_price": fixed_price,
        "product_code": product_code,
    }


def parse_options(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    selects = soup.select("select[option_product_no][product_type='product_option'], select[id^='product_option_id']")
    if not selects:
        return []

    if len(selects) > 1:
        options = []
        for select_index, select in enumerate(selects, start=1):
            title = clean_text(select.get("option_title") or select.get("name") or f"옵션{select_index}")
            for option in select.select("option"):
                label = clean_text(option.get_text(" ", strip=True))
                value = clean_text(option.get("value"))
                if value in ("", "*", "**") or "----" in label or "옵션을 선택" in label:
                    continue
                options.append({
                    "option_title": title,
                    "option_name": label,
                    "option_value": value,
                    "index": len(options),
                    "is_soldout_label": "품절" in label or option.has_attr("disabled"),
                    "multi_select": True,
                })
        return options

    select = selects[0]
    option_title = clean_text(select.get("option_title") or select.get("name") or "옵션")
    options = []
    for option in select.select("option"):
        label = clean_text(option.get_text(" ", strip=True))
        value = clean_text(option.get("value"))
        if value in ("", "*", "**") or "----" in label or "옵션을 선택" in label:
            continue
        options.append({
            "option_title": option_title,
            "option_name": label,
            "option_value": value,
            "index": len(options),
            "is_soldout_label": "품절" in label or option.has_attr("disabled"),
            "multi_select": False,
        })
    return options


def parse_default_item_code(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    selectors = [
        "#totalProducts [item_code]",
        "input[name='option_box_price'][item_code]",
        ".quantity_price [item_code]",
        ".quantity_price[item_code]",
    ]

    for selector in selectors:
        element = soup.select_one(selector)
        item_code = clean_text(element.get("item_code")) if element else ""
        if item_code:
            return item_code

    match = re.search(r'item_code=["\'](P[0-9A-Z]+)["\']', html)
    return match.group(1) if match else ""


def build_calculator_url(product_no: str, cate_no: str, item_code: str, quantity: int = REQUEST_QUANTITY) -> str:
    params = [("product_no", str(product_no)), ("is_subscription", "F"), (f"product[{item_code}]", str(quantity)), ("cate_no", str(cate_no))]
    return CALCULATOR_URL + "?" + urlencode(params)

def get_stock_driver() -> webdriver.Chrome:
    global STOCK_DRIVER
    if STOCK_DRIVER is not None:
        return STOCK_DRIVER

    options = Options()
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-popup-blocking")

    STOCK_DRIVER = webdriver.Chrome(options=options)
    STOCK_DRIVER.set_page_load_timeout(25)
    return STOCK_DRIVER


def close_stock_driver() -> None:
    global STOCK_DRIVER
    if STOCK_DRIVER is None:
        return

    try:
        STOCK_DRIVER.quit()
    except Exception:
        pass

    STOCK_DRIVER = None


def read_url_text_from_browser(url: str) -> str:
    driver = get_stock_driver()
    driver.get(url)
    time.sleep(0.25)

    text = driver.execute_script("return document.body ? document.body.innerText : '';") or ""
    text = text.strip()

    if text:
        return text

    return driver.find_element(By.TAG_NAME, "body").text.strip()


def parse_quantity_alert_limit(message: str) -> int | None:
    patterns = [
        r"최대\s*주문\s*수량은\s*([0-9,]+)\s*개",
        r"최대\s*구매\s*수량은\s*([0-9,]+)\s*개",
        r"([0-9,]+)\s*개까지\s*구매",
    ]

    for pattern in patterns:
        match = re.search(pattern, message or "")
        if match:
            return int(match.group(1).replace(",", ""))

    return None


def accept_current_alert() -> str:
    driver = get_stock_driver()
    try:
        alert = driver.switch_to.alert
        text = clean_text(alert.text)
        alert.accept()
        return text
    except NoAlertPresentException:
        return ""


def get_current_quantity(driver: webdriver.Chrome) -> int:
    try:
        value = driver.execute_script("""
            const input = document.querySelector('#totalProducts input[name="quantity_opt[]"], #totalProducts #quantity, input#quantity');
            return input ? input.value : '';
        """)
        return int(digits_only(str(value)) or "1")
    except Exception:
        return 1


def find_quantity_up_button(driver: webdriver.Chrome):
    selectors = [
        "#totalProducts .QuantityUp",
        "#totalProducts a.up",
        "a.QuantityUp",
        "a.up.QuantityUp",
    ]

    for selector in selectors:
        elements = driver.find_elements(By.CSS_SELECTOR, selector)
        for element in elements:
            try:
                if element.is_displayed():
                    return element
            except Exception:
                continue

    return None


def check_single_product_stock_by_quantity_click(product_url: str) -> dict:
    driver = get_stock_driver()
    request_url = product_url

    try:
        driver.get(product_url)
        time.sleep(0.7)
        accept_current_alert()

        item_code = driver.execute_script("""
            const item = document.querySelector('#totalProducts [item_code], input[name="option_box_price"][item_code], .quantity_price[item_code]');
            return item ? item.getAttribute('item_code') : '';
        """) or ""

        button = find_quantity_up_button(driver)
        if button is None:
            return {
                "stock": "",
                "is_soldout": "N",
                "result": "quantity_button_missing",
                "message": "단일 상품 수량 증가 버튼을 찾지 못했습니다.",
                "request_url": request_url,
                "raw": item_code,
            }

        previous_quantity = get_current_quantity(driver)
        max_clicks = 80

        for _ in range(max_clicks):
            previous_quantity = get_current_quantity(driver)
            driver.execute_script("arguments[0].click();", button)
            time.sleep(0.15)

            alert_message = accept_current_alert()
            if alert_message:
                limit = parse_quantity_alert_limit(alert_message)
                if limit is not None:
                    return {
                        "stock": f"{limit}개 이상(최대주문 {limit})",
                        "is_soldout": "N",
                        "result": "quantity_alert_limit",
                        "message": alert_message,
                        "request_url": request_url,
                        "raw": item_code,
                    }

                return {
                    "stock": str(previous_quantity),
                    "is_soldout": "Y" if previous_quantity <= 0 else "N",
                    "result": "quantity_alert_stock_limit",
                    "message": alert_message,
                    "request_url": request_url,
                    "raw": item_code,
                }

            current_quantity = get_current_quantity(driver)
            if current_quantity <= previous_quantity:
                return {
                    "stock": str(previous_quantity),
                    "is_soldout": "Y" if previous_quantity <= 0 else "N",
                    "result": "quantity_not_increased",
                    "message": "수량 증가 버튼을 눌렀지만 수량이 더 이상 증가하지 않았습니다.",
                    "request_url": request_url,
                    "raw": item_code,
                }

        return {
            "stock": f"{max_clicks + 1}개 이상",
            "is_soldout": "N",
            "result": "quantity_click_max_reached",
            "message": f"{max_clicks}회까지 수량 증가 시도 중 제한 alert가 나오지 않았습니다.",
            "request_url": request_url,
            "raw": item_code,
        }

    except TimeoutException as exc:
        return {
            "stock": "",
            "is_soldout": "N",
            "result": "quantity_page_timeout",
            "message": f"{type(exc).__name__}: {exc}",
            "request_url": request_url,
            "raw": "",
        }
    except Exception as exc:
        return {
            "stock": "",
            "is_soldout": "N",
            "result": "quantity_click_error",
            "message": f"{type(exc).__name__}: {exc}",
            "request_url": request_url,
            "raw": "",
        }

def check_stock(session: requests.Session, product_no: str, cate_no: str, item_code: str, referer_url: str) -> dict:
    """
    EXPED 재고 확인 전용 요청입니다.
    요청 URL을 만든 뒤 실제 브라우저 주소창에 입력한 결과 화면의 텍스트를 수집합니다.
    """
    request_url = build_calculator_url(product_no, cate_no, item_code, REQUEST_QUANTITY)

    try:
        text = read_url_text_from_browser(request_url)

        if not text:
            return {
                "stock": "",
                "is_soldout": "N",
                "result": "empty_browser_body",
                "message": "브라우저 화면에서 응답 텍스트를 읽지 못했습니다.",
                "request_url": request_url,
                "raw": "",
            }

        try:
            data = json.loads(text)
        except Exception as exc:
            return {
                "stock": "",
                "is_soldout": "N",
                "result": "json_error",
                "message": f"{type(exc).__name__}: {exc} / body={text[:300]}",
                "request_url": request_url,
                "raw": text[:1000],
            }

    except Exception as exc:
        return {
            "stock": "",
            "is_soldout": "N",
            "result": "browser_request_error",
            "message": f"{type(exc).__name__}: {exc}",
            "request_url": request_url,
            "raw": "",
        }

    raw = json.dumps(data, ensure_ascii=False)
    message = clean_text(data.get("msg", "")) if isinstance(data, dict) else ""

    if isinstance(data, dict) and "stock_number" in data:
        stock = str(data.get("stock_number", "")).strip()
        return {
            "stock": stock,
            "is_soldout": "Y" if stock == "0" else "N",
            "result": "stock_number",
            "message": message,
            "request_url": request_url,
            "raw": raw,
        }

    if isinstance(data, dict) and item_code in data and isinstance(data.get(item_code), dict):
        return {
            "stock": "",
            "is_soldout": "N",
            "result": "calculated",
            "message": "CalculatorProduct가 stock_number 없이 계산 객체를 반환했습니다. 실제 재고 숫자가 아니므로 수량 클릭 방식으로 재확인해야 합니다.",
            "request_url": request_url,
            "raw": raw,
        }

    if isinstance(data, dict) and data.get("Result") is True:
        return {
            "stock": "",
            "is_soldout": "N",
            "result": "available",
            "message": message or "요청 수량이 통과했지만 정확한 재고 숫자는 응답에 없습니다.",
            "request_url": request_url,
            "raw": raw,
        }

    if isinstance(data, dict) and data.get("Result") is False and "품절" in message:
        return {
            "stock": "0",
            "is_soldout": "Y",
            "result": "soldout_message",
            "message": message,
            "request_url": request_url,
            "raw": raw,
        }

    return {
        "stock": "",
        "is_soldout": "N",
        "result": "unknown",
        "message": message or raw[:500],
        "request_url": request_url,
        "raw": raw,
    }


def discover_candidate_item_codes(session: requests.Session, product_no: str, cate_no: str, product_code: str, product_url: str, visible_count: int) -> list[dict]:
    candidates = []
    max_count = max(visible_count + 8, 26)
    for suffix in option_suffix_candidates(max_count):
        item_code = product_code + suffix
        stock_info = check_stock(session, product_no, cate_no, item_code, product_url)
        if is_useful_stock_response(stock_info):
            candidates.append({
                "suffix": suffix,
                "itemCode": item_code,
                "optionId": suffix,
                "stockInfo": stock_info,
            })
        time.sleep(random.uniform(*SLEEP_BETWEEN_STOCK_REQUESTS))
    return candidates


def collect_product_data(session: requests.Session, product_link: dict) -> dict:
    product_url = product_link["product_url"]
    html = get_text(session, product_url, referer=product_link["list_page_url"])
    meta = parse_product_meta(html, product_link)
    product_no = product_link["product_no"]
    cate_no = product_link["cate_no"]
    product_code = meta["product_code"]
    options = parse_options(html)

    if not product_code:
        return {
            "goodsNo": product_no,
            "productName": meta["product_name"],
            "optionFl": "y" if options else "n",
            "stockCnt": "",
            "goodsPrice": meta["goods_price"],
            "fixedPrice": meta["fixed_price"],
            "options": [],
            "memo": "상품코드 확인 실패",
        }

    if not options:
        item_code = parse_default_item_code(html) or product_code
        stock_info = check_stock(session, product_no, cate_no, item_code, product_url)
        debug_urls = [f"CalculatorProduct: {stock_info.get('request_url', '')}"]
        log(f"    [STOCK] {product_no} 기본상품 item_code={item_code}, qty={REQUEST_QUANTITY} -> {stock_info['stock'] or '미확인'} / {stock_info['result']}")
        log(f"      URL: {stock_info['request_url']}")

        if stock_info.get("result") in ("calculated", "available", "unknown", "empty_browser_body", "json_error", "browser_request_error") or not stock_info.get("stock"):
            click_info = check_single_product_stock_by_quantity_click(product_url)
            debug_urls.append(f"QuantityClick: {click_info.get('request_url', '')}")
            log(f"    [QUANTITY] {product_no} 기본상품 수량 클릭 확인 -> {click_info['stock'] or '미확인'} / {click_info['result']}")
            if click_info.get("message"):
                log(f"      MSG: {click_info.get('message', '')}")
            if click_info.get("stock"):
                stock_info = click_info

        if stock_info.get("result") in ("request_error", "http_error", "json_error", "unknown", "browser_request_error", "quantity_click_error"):
            log(f"      MSG: {stock_info.get('message', '')}")

        return {
            "goodsNo": product_no,
            "productName": meta["product_name"],
            "optionFl": "n",
            "stockCnt": stock_info["stock"],
            "goodsPrice": meta["goods_price"],
            "fixedPrice": meta["fixed_price"],
            "options": [],
            "memo": stock_info["message"],
            "debugUrls": "\n".join(url for url in debug_urls if url.strip()),
        }

    option_rows = []
    direct_results = []
    for option in options:
        if option.get("is_soldout_label"):
            option_rows.append({
                "optionName": clean_text(option["option_name"]).replace(" [품절]", ""),
                "optionTitle": option["option_title"],
                "optionValue": option["option_value"],
                "itemCode": "",
                "optionId": "",
                "optionStock": "0",
                "isSoldOut": "Y",
                "result": "soldout_label",
                "message": "옵션명에 품절 표시가 있어 요청을 생략했습니다.",
                "requestUrl": "옵션명 품절 표시로 요청 생략",
            })
            log(f"    [SKIP] {product_no} {option['option_name']} -> 옵션명 품절 표시")
            continue

        suffix = option_suffix(option["index"])
        item_code = product_code + suffix
        stock_info = check_stock(session, product_no, cate_no, item_code, product_url)
        direct_results.append((option, suffix, item_code, stock_info))
        log(f"    [DIRECT] {product_no} {option['option_name']} item_code={item_code}, qty={REQUEST_QUANTITY} -> {stock_info['stock'] or '미확인'} / {stock_info['result']}")
        log(f"      URL: {stock_info['request_url']}")
        if stock_info.get("result") in ("request_error", "http_error", "json_error", "unknown"):
            log(f"      MSG: {stock_info.get('message', '')}")
        time.sleep(random.uniform(*SLEEP_BETWEEN_STOCK_REQUESTS))

    need_discovery = any(not is_useful_stock_response(item[3]) for item in direct_results)
    discovered = []
    if need_discovery:
        log(f"    [DISCOVER] direct 결과에 미확인이 있어 {product_code}000A~ 후보를 추가 확인합니다.")
        discovered = discover_candidate_item_codes(session, product_no, cate_no, product_code, product_url, len(options))
        log(f"    [DISCOVER] 유효 item_code={len(discovered)}개: {', '.join(d['itemCode'] for d in discovered[:20])}")

    if discovered and len(discovered) >= len(options):
        # Cafe24 옵션코드가 000A부터 연속이 아니거나 삭제된 옵션이 있을 수 있어,
        # 유효 응답이 나온 item_code를 옵션 표시 순서에 다시 매칭합니다.
        final_pairs = []
        for idx, option in enumerate(options):
            d = discovered[idx]
            final_pairs.append((option, d["optionId"], d["itemCode"], d["stockInfo"]))
    else:
        final_pairs = direct_results

    for option, option_id, item_code, stock_info in final_pairs:
        if stock_info["result"] == "request_error" and option.get("multi_select"):
            stock_info["message"] = "다단 옵션 상품일 수 있습니다. 실제 조합 item_code 확인 필요. " + stock_info["message"]

        final_soldout = "Y" if stock_info["is_soldout"] == "Y" or stock_info["stock"] == "0" else "N"
        if stock_info.get("result") in ("calculated", "available"):
            stock_info["stock"] = ""
        option_rows.append({
            "optionName": option["option_name"],
            "optionTitle": option["option_title"],
            "optionValue": option["option_value"],
            "itemCode": item_code,
            "optionId": option_id,
            "optionStock": stock_info["stock"],
            "isSoldOut": final_soldout,
            "result": stock_info["result"],
            "message": stock_info["message"],
            "requestUrl": stock_info["request_url"],
        })
        log(f"    [STOCK] {product_no} {option['option_name']} item_code={item_code}, qty={REQUEST_QUANTITY} -> {stock_info['stock'] or '미확인'} / {stock_info['result']}")

    return {
        "goodsNo": product_no,
        "productName": meta["product_name"],
        "optionFl": "y",
        "stockCnt": "",
        "goodsPrice": meta["goods_price"],
        "fixedPrice": meta["fixed_price"],
        "options": option_rows,
        "memo": "",
        "debugUrls": "\n".join(debug_urls_summary({"options": option_rows})),
    }


def debug_urls_summary(data: dict) -> list[str]:
    urls = []
    for opt in data.get("options") or []:
        name = clean_text(opt.get("optionName"))
        request_url = clean_text(opt.get("requestUrl"))
        if request_url:
            urls.append(f"{name}: {request_url}")
    direct_debug = clean_text(data.get("debugUrls"))
    if direct_debug:
        urls.append(direct_debug)
    return urls


def option_summary(data: dict) -> str:
    lines = []
    for opt in data.get("options") or []:
        name = clean_text(opt.get("optionName"))
        stock = clean_text(opt.get("optionStock"))
        sold = opt.get("isSoldOut") == "Y"
        if not name:
            continue
        if sold or stock == "0":
            lines.append(f"{name} [품절]")
        elif stock:
            lines.append(f"{name} - 재고({stock})")
        else:
            lines.append(f"{name} - 재고(미확인)")
    return "\n".join(lines)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = {
        "A": 12,
        "B": 12,
        "C": 34,
        "D": 10,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 48,
        "I": 22,
        "J": 90,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def create_new_workbook(out_path: str):
    path = Path(out_path)
    if path.exists():
        try:
            path.unlink()
            log(f"[INFO] 기존 파일을 삭제하고 새로 생성합니다: {out_path}")
        except PermissionError:
            raise RuntimeError("선택한 엑셀 파일이 열려 있습니다. 엑셀을 닫고 다시 실행하세요.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(HEADERS)
    style_sheet(ws)
    wb.save(path)
    log("[INFO] new file mode: 기존 행을 건너뛰지 않습니다.")
    return wb, ws, set()


def open_existing_workbook_for_append(out_path: str):
    path = Path(out_path)
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(HEADERS)
        style_sheet(ws)
        wb.save(path)
        log("[INFO] append target did not exist. created new file")
        return wb, ws, set()

    wb = load_workbook(path)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.active
    ws.title = "Products"
    header = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
    expected_old = ["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "optionSummary", "fetchedAt"]
    expected_without_debug = ["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "fixedPrice", "optionSummary", "fetchedAt"]
    if header == HEADERS:
        pass
    elif header == expected_without_debug:
        ws.cell(row=1, column=len(HEADERS)).value = "debugUrls"
        wb.save(path)
    elif header == expected_old:
        ws.insert_cols(7)
        ws.cell(row=1, column=7).value = "fixedPrice"
        for col, name in enumerate(HEADERS, start=1):
            ws.cell(row=1, column=col).value = name
        wb.save(path)
    else:
        raise RuntimeError(f"기존 파일 헤더가 예상과 다릅니다.\n기존: {header}\n예상: {HEADERS}")

    existing = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        goods_no = row[1] if row and len(row) > 1 else None
        if goods_no is not None:
            existing.add(str(goods_no).strip())
    log(f"[INFO] append mode: existing goodsNo={len(existing)}")
    return wb, ws, existing


def append_product_row(ws, data: dict, fetched_at: str) -> None:
    ws.append([
        VENDOR,
        data.get("goodsNo", ""),
        data.get("productName", ""),
        data.get("optionFl", ""),
        data.get("stockCnt", ""),
        data.get("goodsPrice", ""),
        data.get("fixedPrice", ""),
        option_summary(data),
        fetched_at,
        "\n".join(debug_urls_summary(data)),
    ])
    for cell in ws[ws.max_row]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def crawl_to_excel(out_path: str) -> None:
    session = make_session()
    wb, ws, existing = create_new_workbook(out_path)
    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    saved_count = 0
    checked_count = 0
    skipped_count = 0

    log(f"[START] EXPED 전체상품 재고 수집 시작")
    log(f"[CONFIG] CalculatorProduct 요청 수량: {REQUEST_QUANTITY}")
    log(f"[OUTPUT] {out_path}")
    diagnostic = check_stock(session, "302", "28", "P00000LQ000B", "https://www.exped.co.kr/product/%EB%A9%94%EA%B0%80%EC%8A%AC%EB%A6%BD/302/category/28/display/1/")
    log(f"[DIAG] product_no=302, item_code=P00000LQ000B, qty={REQUEST_QUANTITY} -> {diagnostic.get('stock') or '미확인'} / {diagnostic.get('result')}")
    log(f"[DIAG] URL: {diagnostic.get('request_url')}")

    try:
        for product_link in iter_category_products(session):
            goods_no = str(product_link["product_no"])
            if goods_no in existing:
                skipped_count += 1
                continue

            checked_count += 1
            log(f"[PRODUCT {checked_count}] goodsNo={goods_no}, {product_link.get('list_name', '')}")
            try:
                data = collect_product_data(session, product_link)
            except Exception as exc:
                log(f"  [SKIP] goodsNo={goods_no} 수집 실패: {type(exc).__name__}: {exc}")
                continue

            append_product_row(ws, data, fetched_at)
            existing.add(goods_no)
            saved_count += 1

            if saved_count % SAVE_EVERY_PRODUCTS == 0:
                style_sheet(ws)
                wb.save(out_path)
                log(f"[SAVE] saved={saved_count}, checked={checked_count}, skipped_existing={skipped_count} -> {out_path}")

            time.sleep(random.uniform(*SLEEP_BETWEEN_PRODUCTS))

    except KeyboardInterrupt:
        log("[INTERRUPT] 중단 요청. 현재까지 수집된 내용을 저장합니다.")
    finally:
        style_sheet(ws)
        wb.save(out_path)
        log(f"[DONE] 저장 완료: saved={saved_count}, checked={checked_count}, skipped_existing={skipped_count}")
        log(f"[DONE] 파일: {out_path}")


def collect_product_map(session: requests.Session) -> dict[str, dict]:
    product_map = {}
    for link in iter_category_products(session):
        product_map.setdefault(str(link["product_no"]), link)
    return product_map


def update_existing_file(in_path: str, out_path: str) -> None:
    session = make_session()
    wb = load_workbook(in_path)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.active
    header = [cell.value for cell in ws[1]]
    if header != HEADERS:
        expected_without_debug = ["vendor", "goodsNo", "productName", "optionFl", "stockCnt", "goodsPrice", "fixedPrice", "optionSummary", "fetchedAt"]
        if header == expected_without_debug:
            ws.cell(row=1, column=len(HEADERS)).value = "debugUrls"
            header = HEADERS
        else:
            raise RuntimeError(f"엑셀 헤더가 예상과 다릅니다.\n기존: {header}\n예상: {HEADERS}")

    col_index = {name: header.index(name) + 1 for name in header}
    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    log("[START] 기존파일 업데이트 시작")
    log("[INFO] 사이트 전체 상품 링크를 먼저 수집합니다.")
    product_map = collect_product_map(session)
    log(f"[INFO] 현재 사이트 상품 수: {len(product_map)}")

    updated = 0
    checked = 0
    for row_no in range(2, ws.max_row + 1):
        goods_no = str(ws.cell(row_no, col_index["goodsNo"]).value or "").strip()
        if not goods_no:
            continue
        product_link = product_map.get(goods_no)
        if not product_link:
            log(f"[SKIP] goodsNo={goods_no} 현재 카테고리에서 찾지 못했습니다.")
            continue

        checked += 1
        log(f"[UPDATE {checked}] goodsNo={goods_no}")
        try:
            data = collect_product_data(session, product_link)
        except Exception as exc:
            log(f"  [SKIP] goodsNo={goods_no} 업데이트 실패: {type(exc).__name__}: {exc}")
            continue

        ws.cell(row_no, col_index["vendor"]).value = VENDOR
        ws.cell(row_no, col_index["productName"]).value = data.get("productName", "")
        ws.cell(row_no, col_index["optionFl"]).value = data.get("optionFl", "")
        ws.cell(row_no, col_index["stockCnt"]).value = data.get("stockCnt", "")
        ws.cell(row_no, col_index["goodsPrice"]).value = data.get("goodsPrice", "")
        ws.cell(row_no, col_index["fixedPrice"]).value = data.get("fixedPrice", "")
        ws.cell(row_no, col_index["optionSummary"]).value = option_summary(data)
        ws.cell(row_no, col_index["fetchedAt"]).value = fetched_at
        ws.cell(row_no, col_index["debugUrls"]).value = "\n".join(debug_urls_summary(data))
        updated += 1

        if updated % SAVE_EVERY_PRODUCTS == 0:
            style_sheet(ws)
            wb.save(out_path)
            log(f"[SAVE] updated={updated}, checked={checked} -> {out_path}")

        time.sleep(random.uniform(*SLEEP_BETWEEN_PRODUCTS))

    style_sheet(ws)
    wb.save(out_path)
    log(f"[DONE] 업데이트 완료: updated={updated}, checked={checked}")
    log(f"[DONE] 파일: {out_path}")


def run_worker(mode: str, in_path: str | None, out_path: str, root: tk.Tk) -> None:
    try:
        if mode == "new":
            crawl_to_excel(out_path)
        elif mode == "update":
            if not in_path:
                raise RuntimeError("수정할 엑셀 파일이 선택되지 않았습니다.")
            update_existing_file(in_path, out_path)
        root.after(0, lambda: show_popup("완료", f"작업이 완료되었습니다.\n\n{out_path}"))
    except Exception as exc:
        log(f"[ERROR] {type(exc).__name__}: {exc}")
        root.after(0, lambda: show_popup("오류", f"작업 중 오류가 발생했습니다.\n\n{type(exc).__name__}: {exc}", error=True))
    finally:
        close_stock_driver()

def main() -> None:
    mode = ask_mode()
    if mode == "cancel":
        print("[CANCEL] mode select canceled")
        return

    in_path = None
    if mode == "update":
        in_path = ask_open_path()
        if not in_path:
            print("[CANCEL] open canceled")
            return
        default_name = "exped_inventory_updated.xlsx"
    else:
        default_name = "exped_inventory_direct.xlsx"

    out_path = ask_save_path(default_name=default_name)
    if not out_path:
        print("[CANCEL] save canceled")
        return

    global LOG_QUEUE
    root, q = start_log_window()
    LOG_QUEUE = q

    thread = threading.Thread(target=run_worker, args=(mode, in_path, out_path, root), daemon=True)
    thread.start()
    root.mainloop()


if __name__ == "__main__":
    main()
