import html as html_lib
import json
import os
import re
import sys
import time
from datetime import datetime
from typing import Any, Optional
from urllib.parse import parse_qs, urljoin, urlparse, urlunparse, urlencode

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import tkinter as tk
from tkinter import filedialog, messagebox


BASE = "https://www.exped.co.kr"
CATEGORY_URLS = [
    "https://www.exped.co.kr/product/list.html?cate_no=24",
    "https://www.exped.co.kr/product/list.html?cate_no=25",
    "https://www.exped.co.kr/product/list.html?cate_no=26",
    "https://www.exped.co.kr/product/list.html?cate_no=27",
    "https://www.exped.co.kr/product/list.html?cate_no=42",
    "https://www.exped.co.kr/product/list.html?cate_no=28",
    "https://www.exped.co.kr/product/list.html?cate_no=43",
    "https://www.exped.co.kr/product/list.html?cate_no=44",
]

HEADERS = ["product_no", "product_name", "sale_price", "consumer_price", "brand", "option_summary", "detail_url", "category_url", "fetched_at", "memo"]
SAVE_EVERY_PRODUCTS = 10
REQUEST_SLEEP_SEC = 0.35
MAX_EMPTY_PAGE_COUNT = 2


try:
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
except Exception:
    pass


SESSION = requests.Session()
SESSION.headers.update({
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
    "Connection": "keep-alive",
})


def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def digits_only(value: str | None) -> str:
    return re.sub(r"[^0-9]", "", value or "")


def fetch(url: str, referer: str | None = None) -> str:
    response = SESSION.get(url, headers={"Referer": referer or BASE}, timeout=30)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or "utf-8"
    return response.text


def absolute_url(href: str | None) -> str:
    return urljoin(BASE, href or "").split("#")[0]


def extract_product_no(value: str | None) -> str:
    text = value or ""
    patterns = [r"product_no=(\d+)", r"/product/[^/]+/(\d+)/", r"anchorBoxId_(\d+)"]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return ""


def extract_cate_no(value: str | None) -> str:
    text = value or ""
    patterns = [r"cate_no=(\d+)", r"/category/(?:[^/]+/)?(\d+)/", r"/category/[^/]+/(\d+)/"]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return ""


def make_page_url(category_url: str, page: int) -> str:
    parsed = urlparse(category_url)
    query = parse_qs(parsed.query)
    query["page"] = [str(page)]
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, urlencode(query, doseq=True), ""))


def parse_last_page(html: str) -> int:
    soup = BeautifulSoup(html, "html.parser")
    pages = [1]

    for a in soup.select(".xans-product-normalpaging a[href*='page=']"):
        href = a.get("href") or ""
        page = parse_qs(urlparse(href).query).get("page", [""])[0]
        if page.isdigit():
            pages.append(int(page))

    last_link = soup.select_one(".xans-product-normalpaging p.last a[href*='page=']")
    if last_link:
        page = parse_qs(urlparse(last_link.get("href", "")).query).get("page", [""])[0]
        if page.isdigit():
            pages.append(int(page))

    return max(pages)


def parse_price_to_int(text: str | None) -> str:
    return digits_only(clean_text(text))


def read_labeled_list_item(li, labels: tuple[str, ...]) -> str:
    for item_li in li.select("ul.xans-product-listitem > li, .spec > li, .description li"):
        label_el = item_li.select_one("strong span, .title span, .title")
        value_el = item_li.select_one("strong ~ span, .title + span, span:not(.title)")
        label = clean_text(label_el.get_text(" ", strip=True)) if label_el else ""
        value = clean_text(value_el.get_text(" ", strip=True)) if value_el else ""
        if any(target in label for target in labels):
            return value
    return ""


def parse_list_product_name(li) -> str:
    selectors = [
        "p.name > a strong.title + span",
        "strong.name a span:not(.title) span",
        "strong.name a span:not(.title)",
        "p.name > a > span:last-child",
        "strong.name a",
        "p.name a",
    ]

    for selector in selectors:
        element = li.select_one(selector)
        text = clean_text(element.get_text(" ", strip=True)) if element else ""
        text = re.sub(r"^상품명\s*", "", text).strip()
        if text and text != "상품명":
            return text

    img = li.select_one(".prdImg img, img")
    return clean_text(img.get("alt")) if img else ""


def parse_products_from_list(html: str, category_url: str, list_page_url: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    root = soup.select_one("#contents") or soup
    products: list[dict[str, str]] = []
    seen: set[str] = set()

    for li in root.select("ul.prdList > li[id^='anchorBoxId_']"):
        link = li.select_one(".prdImg a[href*='/product/']") or li.select_one("p.name a[href*='/product/']") or li.select_one("strong.name a[href*='/product/']") or li.select_one("a[href*='/product/']")
        if not link:
            continue

        detail_url = absolute_url(link.get("href"))
        product_no = extract_product_no(detail_url) or extract_product_no(li.get("id"))
        cate_no = extract_cate_no(detail_url) or extract_cate_no(category_url)
        if not product_no or product_no in seen:
            continue

        sale_price = ""
        price_node = li.select_one(".description[ec-data-price]")
        if price_node:
            sale_price = parse_price_to_int(price_node.get("ec-data-price"))
        if not sale_price:
            sale_price = parse_price_to_int(read_labeled_list_item(li, ("판매가",)))

        consumer_price = parse_price_to_int(read_labeled_list_item(li, ("소비자가", "정가")))
        brand = read_labeled_list_item(li, ("브랜드",))

        products.append({
            "product_no": product_no,
            "cate_no": cate_no,
            "product_name": parse_list_product_name(li),
            "sale_price": sale_price,
            "consumer_price": consumer_price,
            "brand": brand,
            "detail_url": detail_url,
            "category_url": category_url,
            "list_page_url": list_page_url,
        })
        seen.add(product_no)

    return products


def extract_product_name(detail_html: str, fallback: str = "") -> str:
    soup = BeautifulSoup(detail_html, "html.parser")
    selectors = [
        ".infoArea .headingArea h1",
        "div.detailArea h3",
        "div.detailArea h2",
        "h1",
        "h2",
    ]

    for selector in selectors:
        element = soup.select_one(selector)
        if not element:
            continue
        for hidden in element.select(".displaynone"):
            hidden.decompose()
        text = clean_text(element.get_text(" ", strip=True))
        if text:
            return text

    title = clean_text(soup.title.get_text(" ", strip=True)) if soup.title else ""
    return title or fallback


def extract_detail_prices(detail_html: str) -> tuple[str, str]:
    soup = BeautifulSoup(detail_html, "html.parser")
    sale_price = ""
    consumer_price = ""

    price_el = soup.select_one("#span_product_price_text") or soup.select_one(".price-original")
    if price_el:
        sale_price = parse_price_to_int(price_el.get_text(" ", strip=True))

    custom_el = soup.select_one("#ec-product-price-info[ec-data-custom]")
    if custom_el:
        custom_value = parse_price_to_int(custom_el.get("ec-data-custom"))
        if custom_value and custom_value != "0":
            consumer_price = custom_value

    for row in soup.select(".infoArea table tr, .xans-product-detaildesign table tr"):
        label = clean_text(row.select_one("th, .title").get_text(" ", strip=True)) if row.select_one("th, .title") else ""
        value = clean_text(row.select_one("td").get_text(" ", strip=True)) if row.select_one("td") else ""
        if not sale_price and "판매가" in label:
            sale_price = parse_price_to_int(value)
        if not consumer_price and ("소비자가" in label or "정가" in label):
            consumer_price = parse_price_to_int(value)

    return sale_price, consumer_price


def load_json_candidate(raw: str) -> dict[str, Any]:
    raw = html_lib.unescape(raw or "").replace(r"\/", "/").strip()
    candidates = [raw]

    try:
        candidates.append(raw.encode("utf-8").decode("unicode_escape"))
    except Exception:
        pass

    for candidate in candidates:
        try:
            data = json.loads(candidate)
            if isinstance(data, dict):
                return data
        except Exception:
            continue

    return {}


def extract_stock_json(detail_html: str, variable_name: str) -> dict[str, Any]:
    patterns = [
        rf"{re.escape(variable_name)}\s*=\s*'([^']*)'\s*;",
        rf'{re.escape(variable_name)}\s*=\s*"([^"]*)"\s*;',
    ]

    for pattern in patterns:
        match = re.search(pattern, detail_html)
        if not match:
            continue
        data = load_json_candidate(match.group(1))
        if data:
            return data

    return {}


def extract_option_stock_data(detail_html: str) -> dict[str, Any]:
    return extract_stock_json(detail_html, "option_stock_data")


def extract_single_stock(detail_html: str) -> tuple[str, bool, str]:
    single_data = extract_stock_json(detail_html, "single_option_stock_data")
    if single_data:
        stock = clean_text(str(single_data.get("stock_number", "")))
        is_selling = clean_text(str(single_data.get("is_selling", "")))
        is_display = clean_text(str(single_data.get("is_display", "")))
        soldout = stock == "0" or (is_selling and is_selling != "T") or (is_display and is_display != "T")
        return stock, soldout, "single_option_stock_data"

    match = re.search(r"var\s+stock_number\s*=\s*['\"](\d+)['\"]\s*;", detail_html)
    if match:
        stock = match.group(1).strip()
        return stock, stock == "0", "stock_number_variable"

    if re.search(r"품절|SOLD\s*OUT", detail_html, re.IGNORECASE):
        return "0", True, "soldout_text"

    return "", False, "stock_not_found"


def parse_visible_options(detail_html: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(detail_html, "html.parser")
    options = []

    for select in soup.select("select[option_product_no][product_type='product_option'], select[id^='product_option_id']"):
        option_title = clean_text(select.get("option_title") or select.get("name") or "옵션")
        for option in select.select("option"):
            value = clean_text(option.get("value"))
            name = clean_text(option.get_text(" ", strip=True))
            if value in ("", "*", "**") or "----" in name or "옵션을 선택" in name:
                continue
            options.append({
                "option_title": option_title,
                "option_name": name.replace("[품절]", "").strip(),
                "is_soldout": "Y" if "품절" in name or option.has_attr("disabled") else "N",
            })

    return options


def build_option_summary(option_data: dict[str, Any], detail_html: str = "") -> tuple[str, str]:
    if option_data:
        lines: list[str] = []
        sortable: list[tuple[str, str, dict[str, Any]]] = []

        for item_code, value in option_data.items():
            if not isinstance(value, dict):
                continue
            option_value = clean_text(str(value.get("option_value", "")))
            sortable.append((option_value, str(item_code), value))

        sortable.sort(key=lambda item: (item[0] == "", item[0], item[1]))

        for option_value, item_code, value in sortable:
            name = option_value or item_code
            stock_raw = value.get("stock_number", "")
            is_selling = clean_text(str(value.get("is_selling", "")))
            is_display = clean_text(str(value.get("is_display", "")))
            stock = clean_text(str(stock_raw))
            stock_int: Optional[int] = None

            try:
                stock_int = int(stock)
            except Exception:
                stock_int = None

            soldout = stock_int == 0 or (is_selling and is_selling != "T") or (is_display and is_display != "T")
            if soldout:
                lines.append(f"{name} [품절]")
            elif stock:
                lines.append(f"{name} - 재고({stock})")
            else:
                lines.append(f"{name} - 재고(미확인)")

        return "\n".join(lines).strip(), "option_stock_data"

    visible_options = parse_visible_options(detail_html)
    if visible_options:
        lines = []
        for option in visible_options:
            if option.get("is_soldout") == "Y":
                lines.append(f"{option['option_name']} [품절]")
            else:
                lines.append(f"{option['option_name']} - 재고(미확인)")
        return "\n".join(lines).strip(), "visible_option_label"

    return "", "option_stock_data_not_found"


def ask_save_xlsx_path() -> str:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    path = filedialog.asksaveasfilename(
        title="EXPED 재고 엑셀(.xlsx) 저장 위치 선택",
        defaultextension=".xlsx",
        filetypes=[("Excel 파일", "*.xlsx")],
        initialfile="exped_inventory_bpl_style.xlsx",
    )
    root.destroy()
    return path


def show_info(title: str, message: str) -> None:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    messagebox.showinfo(title, message)
    root.destroy()


def show_error(title: str, message: str) -> None:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    messagebox.showerror(title, message)
    root.destroy()


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    widths = {
        "A": 12,
        "B": 46,
        "C": 14,
        "D": 14,
        "E": 20,
        "F": 62,
        "G": 82,
        "H": 42,
        "I": 22,
        "J": 32,
    }

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def open_or_create_workbook(out_path: str) -> tuple[Workbook, Any, set[str]]:
    if os.path.exists(out_path):
        wb = load_workbook(out_path)
        ws = wb["Products"] if "Products" in wb.sheetnames else wb.active
        ws.title = "Products"
        header = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
        if header != HEADERS:
            raise RuntimeError(f"기존 파일 헤더가 예상과 다릅니다.\n기존: {header}\n예상: {HEADERS}")

        existing = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                existing.add(str(row[0]).strip())
        return wb, ws, existing

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(HEADERS)
    style_sheet(ws)
    return wb, ws, set()


def collect_detail_product(product: dict[str, str]) -> dict[str, str]:
    detail_html = fetch(product["detail_url"], referer=product["list_page_url"])
    detail_sale_price, detail_consumer_price = extract_detail_prices(detail_html)
    product_name = product.get("product_name") or extract_product_name(detail_html)
    option_data = extract_option_stock_data(detail_html)
    option_summary, memo = build_option_summary(option_data, detail_html)

    if not option_summary:
        single_stock, single_soldout, single_memo = extract_single_stock(detail_html)
        memo = single_memo
        if single_stock:
            option_summary = "단품 [품절]" if single_soldout else f"단품 - 재고({single_stock})"

    if not option_summary:
        option_summary = "재고(미확인)"

    return {
        "product_no": product["product_no"],
        "product_name": product_name,
        "sale_price": product.get("sale_price") or detail_sale_price,
        "consumer_price": product.get("consumer_price") or detail_consumer_price,
        "brand": product.get("brand", ""),
        "option_summary": option_summary,
        "detail_url": product["detail_url"],
        "category_url": product["category_url"],
        "memo": memo,
    }


def append_product_row(ws, data: dict[str, str], fetched_at: str) -> None:
    ws.append([
        data.get("product_no", ""),
        data.get("product_name", ""),
        data.get("sale_price", ""),
        data.get("consumer_price", ""),
        data.get("brand", ""),
        data.get("option_summary", ""),
        data.get("detail_url", ""),
        data.get("category_url", ""),
        fetched_at,
        data.get("memo", ""),
    ])

    for cell in ws[ws.max_row]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def crawl_to_excel(out_path: str) -> None:
    wb, ws, existing_products = open_or_create_workbook(out_path)
    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total_added = 0
    total_seen = 0

    print("[START] EXPED 재고 수집 시작")
    print(f"[OUTPUT] {out_path}")

    for category_index, category_url in enumerate(CATEGORY_URLS, start=1):
        print(f"[CATEGORY {category_index}/{len(CATEGORY_URLS)}] {category_url}")
        try:
            first_html = fetch(category_url)
        except Exception as exc:
            print(f"  [SKIP] 카테고리 접근 실패: {type(exc).__name__}: {exc}")
            continue

        last_page = parse_last_page(first_html)
        empty_count = 0

        for page in range(1, last_page + 1):
            page_url = make_page_url(category_url, page)
            try:
                list_html = first_html if page == 1 else fetch(page_url, referer=category_url)
            except Exception as exc:
                print(f"  [SKIP] page={page} 접근 실패: {type(exc).__name__}: {exc}")
                continue

            products = parse_products_from_list(list_html, category_url, page_url)
            print(f"  [LIST] page={page}/{last_page}, 상품={len(products)}")

            if not products:
                empty_count += 1
                if empty_count >= MAX_EMPTY_PAGE_COUNT:
                    print("  [STOP] 빈 페이지가 반복되어 해당 카테고리를 종료합니다.")
                    break
                continue

            empty_count = 0

            for product in products:
                product_no = product["product_no"]
                total_seen += 1
                if product_no in existing_products:
                    continue

                try:
                    data = collect_detail_product(product)
                except Exception as exc:
                    print(f"    [SKIP] product_no={product_no} 상세 수집 실패: {type(exc).__name__}: {exc}")
                    continue

                append_product_row(ws, data, fetched_at)
                existing_products.add(product_no)
                total_added += 1
                print(f"    [ADD] {product_no} / {data.get('product_name', '')} / {data.get('memo', '')}")

                if total_added % SAVE_EVERY_PRODUCTS == 0:
                    style_sheet(ws)
                    wb.save(out_path)
                    print(f"[SAVE] added={total_added}, seen={total_seen} -> {out_path}")

                time.sleep(REQUEST_SLEEP_SEC)

    style_sheet(ws)
    wb.save(out_path)
    print(f"[DONE] 저장 완료: added={total_added}, seen={total_seen}")
    print(f"[DONE] 파일: {out_path}")


def main() -> None:
    out_path = ask_save_xlsx_path()
    if not out_path:
        show_info("취소됨", "저장 위치 선택이 취소되어 작업을 종료합니다.")
        return

    try:
        crawl_to_excel(out_path)
        show_info("완료", f"재고 수집이 완료되었습니다.\n\n{out_path}")
    except Exception as exc:
        show_error("오류", f"작업 중 오류가 발생했습니다.\n\n{type(exc).__name__}: {exc}")
        raise


if __name__ == "__main__":
    main()
