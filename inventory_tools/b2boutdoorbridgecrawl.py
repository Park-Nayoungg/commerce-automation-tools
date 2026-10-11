from __future__ import annotations

from pathlib import Path as _ConfigPath
import sys as _config_sys
_config_sys.path.insert(0, str(_ConfigPath(__file__).resolve().parents[1]))
from commerce_config import configured_url, output_dir
_OUTDOOR_START_URL = configured_url('outdoor_start_url')


import re
import time
from pathlib import Path
import threading
from tkinter import Tk, filedialog, Button, Label
from urllib.parse import urljoin

import pandas as pd
from openpyxl.styles import Alignment
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


START_URL = _OUTDOOR_START_URL
BASE_URL = "https://www.b2boutdoor.co.kr"
WAIT_TIMEOUT = 15


def ask_output_file() -> str:
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    output_file = filedialog.asksaveasfilename(
        title="재고 엑셀 파일 저장 위치를 선택하세요.",
        defaultextension=".xlsx",
        initialfile="b2b_products_stock.xlsx",
        filetypes=[("Excel 파일", "*.xlsx")]
    )

    root.destroy()
    return output_file

class StopPopup:
    def __init__(self) -> None:
        self.stop_event = threading.Event()
        self.thread = threading.Thread(target=self._run_popup, daemon=True)

    def start(self) -> None:
        self.thread.start()

    def should_stop(self) -> bool:
        return self.stop_event.is_set()

    def _run_popup(self) -> None:
        root = Tk()
        root.title("재고 수집 제어")
        root.geometry("280x130")
        root.resizable(False, False)
        root.attributes("-topmost", True)

        label = Label(
            root,
            text="수집을 중단하려면\n아래 버튼을 누르세요.",
            font=("맑은 고딕", 10),
            justify="center"
        )
        label.pack(pady=(18, 10))

        stop_button = Button(
            root,
            text="수집 종료",
            width=16,
            height=2,
            command=lambda: self._request_stop(root)
        )
        stop_button.pack()

        root.protocol("WM_DELETE_WINDOW", lambda: self._request_stop(root))
        root.mainloop()

    def _request_stop(self, root: Tk) -> None:
        self.stop_event.set()
        root.destroy()

def create_driver() -> webdriver.Chrome:
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_experimental_option("detach", True)

    service = Service()
    driver = webdriver.Chrome(service=service, options=options)
    return driver


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def clean_price(text: str) -> str:
    return normalize_text(text)


def wait_for_product_list(driver: webdriver.Chrome) -> None:
    WebDriverWait(driver, WAIT_TIMEOUT).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, "div.product_list"))
    )


def wait_for_detail_page(driver: webdriver.Chrome) -> None:
    WebDriverWait(driver, WAIT_TIMEOUT).until(
        EC.presence_of_element_located((By.ID, "product_history_back"))
    )


def parse_detail_click_args(href: str) -> tuple[str, str, str] | None:
    match = re.search(
        r"product_detail_click\('([^']+)'\s*,\s*'([^']+)'\s*,\s*'([^']+)'\)",
        href or ""
    )

    if not match:
        return None

    comp_cd, brand_code, product_code = match.groups()
    return comp_cd, brand_code, product_code


def crawl_current_page_products(driver: webdriver.Chrome, page_no: int) -> list[dict]:
    wait_for_product_list(driver)

    items = driver.find_elements(By.CSS_SELECTOR, "div.product_list")
    products: list[dict] = []

    for item in items:
        try:
            anchors = item.find_elements(By.CSS_SELECTOR, "dl dd a")
            img = item.find_element(By.CSS_SELECTOR, "dl dd img")
            dds = item.find_elements(By.CSS_SELECTOR, "dl dd")

            if len(anchors) < 3 or len(dds) < 4:
                continue

            detail_href = anchors[1].get_attribute("href") or anchors[1].get_attribute("outerHTML") or ""
            detail_args = parse_detail_click_args(detail_href)

            if detail_args is None:
                print("[경고] 상세 페이지 이동 함수 파싱 실패")
                continue

            comp_cd, brand_code, product_code = detail_args
            product_name = normalize_text(anchors[2].text)
            price = clean_price(dds[3].text)
            image_src = img.get_attribute("src") or ""

            if image_src.startswith("/"):
                image_src = urljoin(BASE_URL, image_src)

            products.append({
                "page": page_no,
                "product_name": product_name,
                "price": price,
                "comp_cd": comp_cd,
                "brand_code": brand_code,
                "product_code": product_code,
                "image_url": image_src
            })

        except Exception as e:
            print(f"[경고] 상품 목록 1개 파싱 실패: {e}")

    return products


def open_product_detail(driver: webdriver.Chrome, product: dict) -> None:
    driver.execute_script(
        "product_detail_click(arguments[0], arguments[1], arguments[2]);",
        product["comp_cd"],
        product["brand_code"],
        product["product_code"]
    )
    wait_for_detail_page(driver)
    time.sleep(0.5)


def is_out_of_stock(stock_text: str) -> bool:
    normalized = normalize_text(stock_text)
    return any(mark in normalized for mark in ["Ⅹ", "X", "×", "x"])


def format_option_text(color: str, size: str, out_of_stock: bool) -> str:
    parts: list[str] = []

    if color:
        parts.append(f"{color}(색상)")

    if size:
        parts.append(f"{size}(사이즈)")

    if parts:
        option_text = " / ".join(parts)
        return f"{option_text} [품절]" if out_of_stock else option_text

    return "[품절]" if out_of_stock else ""

def extract_stock_text_from_detail(driver: webdriver.Chrome) -> str:
    stock_tables = driver.find_elements(
        By.XPATH,
        "//form//table[.//tr/td[normalize-space()='재고여부']]"
    )

    if not stock_tables:
        print("[경고] 재고 테이블을 찾지 못했습니다.")
        return "재고 테이블 확인 필요"

    stock_table = stock_tables[0]
    rows = stock_table.find_elements(By.XPATH, "./tbody/tr")
    option_lines: list[str] = []

    for row in rows:
        tds = row.find_elements(By.XPATH, "./td")

        if len(tds) < 3:
            continue

        row_text = normalize_text(row.text)

        if "재고여부" in row_text:
            continue

        if len(tds) >= 6:
            color = normalize_text(tds[0].text)
            size = normalize_text(tds[1].text)
            stock_text = normalize_text(tds[2].text)
        else:
            continue

        if not stock_text:
            continue

        out_of_stock = is_out_of_stock(stock_text)
        option_text = format_option_text(color, size, out_of_stock)

        if option_text:
            option_lines.append(option_text)

    return "\n".join(option_lines)


def return_to_product_list(driver: webdriver.Chrome) -> None:
    back_button = WebDriverWait(driver, WAIT_TIMEOUT).until(
        EC.element_to_be_clickable((By.ID, "product_history_back"))
    )

    driver.execute_script("arguments[0].click();", back_button)
    wait_for_product_list(driver)
    time.sleep(0.5)


def collect_product_stock(driver: webdriver.Chrome, product: dict) -> str:
    try:
        open_product_detail(driver, product)
        stock_text = extract_stock_text_from_detail(driver)
        return_to_product_list(driver)
        return stock_text

    except Exception as e:
        print(f"[경고] 상세 재고 확인 실패: {product['product_name']} / {e}")

        try:
            back_buttons = driver.find_elements(By.ID, "product_history_back")

            if back_buttons:
                driver.execute_script("arguments[0].click();", back_buttons[0])
                wait_for_product_list(driver)

        except Exception as back_error:
            print(f"[경고] 목록 복귀 실패: {back_error}")

        return "상세 조회 실패"


def get_current_page_number(driver: webdriver.Chrome) -> int:
    try:
        current_font = driver.find_element(By.CSS_SELECTOR, "p font.font_bold")
        current_text = normalize_text(current_font.text)
        return int(current_text)
    except Exception:
        return 1


def go_to_next_page(driver: webdriver.Chrome) -> bool:
    try:
        current_page = get_current_page_number(driver)
        page_links = driver.find_elements(By.CSS_SELECTOR, "p a[href*='product_page_click']")
        next_link = None

        for link in page_links:
            text = normalize_text(link.text)

            if text.isdigit() and int(text) == current_page + 1:
                next_link = link
                break

        if next_link is None:
            return False

        driver.execute_script("arguments[0].click();", next_link)

        WebDriverWait(driver, WAIT_TIMEOUT).until(
            lambda d: get_current_page_number(d) == current_page + 1
        )

        wait_for_product_list(driver)
        time.sleep(1)
        return True

    except Exception as e:
        print(f"[경고] 다음 페이지 이동 실패: {e}")
        return False


def save_to_excel(data: list[dict], output_file: str) -> None:
    df = pd.DataFrame(data, columns=["상품명", "판매가", "옵션 및 재고여부"])

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="재고현황")
        worksheet = writer.book["재고현황"]

        worksheet.column_dimensions["A"].width = 45
        worksheet.column_dimensions["B"].width = 18
        worksheet.column_dimensions["C"].width = 60

        for row in worksheet.iter_rows(min_row=1):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for row_index in range(2, worksheet.max_row + 1):
            option_value = worksheet.cell(row=row_index, column=3).value or ""
            line_count = str(option_value).count("\n") + 1
            worksheet.row_dimensions[row_index].height = max(18, line_count * 16)

    print(f"[완료] 엑셀 저장: {output_file}")

def main() -> None:
    output_file = ask_output_file()

    if not output_file:
        print("[종료] 엑셀 저장 위치가 선택되지 않았습니다.")
        return

    stop_popup = StopPopup()
    stop_popup.start()

    driver = create_driver()
    final_results: list[dict] = []

    try:
        driver.get(START_URL)
        print("1단계: 브라우저가 열렸습니다.")
        input("직접 로그인 완료 후 Enter를 누르세요: ")

        print("2단계: 사이트에서 검색 조건을 직접 설정하고 검색 버튼까지 누르세요.")
        input("검색 결과 페이지가 화면에 보이면 Enter를 누르세요: ")

        wait_for_product_list(driver)

        while True:
            if stop_popup.should_stop():
                print("[중단] 사용자가 수집 종료를 요청했습니다.")
                break

            current_page = get_current_page_number(driver)
            print(f"\n[수집 시작] {current_page}페이지")

            products = crawl_current_page_products(driver, current_page)
            print(f"[상품 목록 확인] {len(products)}개")

            for index, product in enumerate(products, start=1):
                if stop_popup.should_stop():
                    print("[중단] 사용자가 수집 종료를 요청했습니다.")
                    break

                print(f"[재고 확인 중] {current_page}페이지 {index}/{len(products)} - {product['product_name']}")

                stock_text = collect_product_stock(driver, product)

                final_results.append({
                    "상품명": product["product_name"],
                    "판매가": product["price"],
                    "옵션 및 재고여부": stock_text
                })

            if stop_popup.should_stop():
                break

            moved = go_to_next_page(driver)

            if not moved:
                print("[종료] 마지막 페이지까지 재고 확인을 완료했습니다.")
                break

        if final_results:
            save_to_excel(final_results, output_file)
        else:
            print("[종료] 저장할 수집 데이터가 없습니다.")

    finally:
        print("브라우저는 열어둔 상태입니다. 직접 닫으시면 됩니다.")

if __name__ == "__main__":
    main()
