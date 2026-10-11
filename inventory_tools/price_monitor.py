
from pathlib import Path as _ConfigPath
import sys as _config_sys
_config_sys.path.insert(0, str(_ConfigPath(__file__).resolve().parents[1]))
from commerce_config import configured_url, output_dir
_GODOMALL_BASE_URL = configured_url('godomall_base_url')

import re
import threading
import queue
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service

from webdriver_manager.chrome import ChromeDriverManager


# =========================
# 설정 / 런타임 입력값
# =========================
@dataclass
class Config:
    admin_url: str = _GODOMALL_BASE_URL + '/design/design_skin_list.php'
    goods_list_url: str = _GODOMALL_BASE_URL + '/goods/goods_list.php'

    xlsx_path: str = ""
    sheet_name: str = "Sheet1"
    start_row: int = 3

    code_col: str | None = None     # 품번(없으면 None)
    name_col: str = ""              # 상품명/검색키워드(필수)
    excel_price_col: str = ""       # 비교할 가격(필수)

    row_fill_cols: tuple[str, str] = ("A", "F")

    # goods_list 검색창/요약
    search_input_xpath: str = "/html/body/div[1]/div/div[2]/div[2]/div/form[1]/div[2]/table/tbody[1]/tr[2]/td/div/input"
    summary_div_xpath: str = "/html/body/div[1]/div/div[2]/div[2]/div/form[1]/div[4]/div[1]"

    # 검색 결과 링크(상품명 a)
    result_name_links_css: str = "form table tbody a.text-blue.hand"

    # 팝업(상품 등록/수정) 정가 input
    fixed_price_css: str = "input[name='fixedPrice']"


# 색상
FILL_PINK = PatternFill(start_color="FFFFC0CB", end_color="FFFFC0CB", fill_type="solid")     # 미검색/에러(행)
FILL_YELLOW = PatternFill(start_color="FFFFF59D", end_color="FFFFF59D", fill_type="solid")  # 가격 불일치(가격 셀)
FILL_GREEN = PatternFill(start_color="FFC8E6C9", end_color="FFC8E6C9", fill_type="solid")   # 가격 일치(행)

# 전역 종료 플래그
stop_flag = {"stop": False}

SNAPSHOT_PATH = "parkers_watch_snapshot.json"
WATCH_KEYWORDS = ["트란지아", "힐레베르그"]
POLL_SECONDS = 180  # 3분(원하시면 60~300으로)

# =========================
# UI: 중지 버튼 팝업
# =========================
def run_stop_popup():
    """
    항상 떠있는 '중지' 버튼 팝업.
    버튼을 누르면 stop_flag["stop"] = True
    """
    import tkinter as tk

    root = tk.Tk()
    root.title("Price Monitor - 실행중")
    root.attributes("-topmost", True)
    root.geometry("260x120")

    label = tk.Label(root, text="실행중입니다.\n중지하려면 아래 버튼을 누르세요.")
    label.pack(pady=10)

    def on_stop():
        stop_flag["stop"] = True
        label.config(text="중지 요청됨\n현재 작업 완료 후 종료합니다.")
        btn.config(state="disabled")

    btn = tk.Button(root, text="중지", width=12, command=on_stop)
    btn.pack(pady=5)

    root.mainloop()


# =========================
# 입력: 엑셀 파일 선택 + 컬럼 지정
# =========================
def pick_excel_file() -> str:
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="모니터링할 엑셀 파일을 선택하세요",
        filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")]
    )
    if not path:
        raise SystemExit("엑셀 파일을 선택하지 않아 종료합니다.")
    return path


def ask_col(prompt: str, required: bool) -> str | None:
    while True:
        v = input(prompt).strip().upper()
        if v == "":
            if required:
                print("필수 입력입니다. 예: D")
                continue
            return None
        if re.fullmatch(r"[A-Z]{1,3}", v):
            return v
        print("열 문자가 올바르지 않습니다. 예: A, C, D, AA")


def build_cfg_from_user_input() -> Config:
    xlsx_path = pick_excel_file()

    print("\n엑셀 컬럼을 입력하세요. 예: A, C, D, AA")
    code_col = ask_col("품번 열(없으면 Enter): ", required=False)
    name_col = ask_col("상품명/검색키워드 열(필수): ", required=True)
    price_col = ask_col("비교할 가격 열(필수): ", required=True)

    start_row_raw = input("데이터 시작 행 번호(기본 3): ").strip()
    start_row = int(start_row_raw) if start_row_raw.isdigit() else 3

    sheet_name = input("시트명(기본 Sheet1): ").strip() or "Sheet1"

    row_start = (input("행 색칠 시작 열(기본 A): ").strip().upper() or "A")
    row_end = (input("행 색칠 끝 열(기본 F): ").strip().upper() or "F")

    return Config(
        xlsx_path=xlsx_path,
        sheet_name=sheet_name,
        start_row=start_row,
        code_col=code_col,
        name_col=name_col,
        excel_price_col=price_col,
        row_fill_cols=(row_start, row_end),
    )


# =========================
# Selenium 유틸
# =========================
def build_driver() -> webdriver.Chrome:
    options = webdriver.ChromeOptions()
    options.add_argument("--window-size=1400,900")

    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)

    driver.implicitly_wait(1)
    return driver


def wait_user_login(driver: webdriver.Chrome, cfg: Config):
    driver.get(cfg.admin_url)
    driver.execute_script("alert('고도몰 관리자에 로그인 완료하신 뒤, 확인을 눌러주세요.');")


# =========================
# Excel 유틸
# =========================
def col_to_idx(col: str) -> int:
    return ord(col.strip().upper()) - ord("A") + 1


def fill_row_range(ws, row: int, start_col: str, end_col: str, fill: PatternFill | None):
    s = col_to_idx(start_col)
    e = col_to_idx(end_col)
    if s > e:
        s, e = e, s
    for col_idx in range(s, e + 1):
        ws.cell(row=row, column=col_idx).fill = fill if fill is not None else PatternFill()


# =========================
# 파싱 유틸
# =========================
def normalize(s: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", s or "").upper()


def to_int_price(v) -> int | None:
    if v is None:
        return None
    s = re.sub(r"[^\d]", "", str(v))
    return int(s) if s else None


# =========================
# 스냅샷 비교 유틸
# =========================












# =========================
# 팝업 로그
# =========================






# =========================
# 고도몰 로직
# =========================
def search_in_goods_list(driver: webdriver.Chrome, cfg: Config, keyword: str) -> bool:
    """
    return True  => 검색결과 0개
    return False => 0개 아님
    """
    wait = WebDriverWait(driver, 15)
    driver.get(cfg.goods_list_url)

    inp = wait.until(EC.presence_of_element_located((By.XPATH, cfg.search_input_xpath)))
    inp.clear()
    inp.send_keys(keyword)
    inp.send_keys(Keys.ENTER)

    summary = wait.until(EC.presence_of_element_located((By.XPATH, cfg.summary_div_xpath)))
    text = (summary.text or "").replace("\n", " ").strip()
    return "검색 0개" in text


def pick_product_link(driver: webdriver.Chrome, cfg: Config, code: str, name_hint: str):
    """
    1) 품번 포함 링크 우선
    2) 없으면 상품명 토큰 포함 점수로 가장 가까운 1개
    """
    wait = WebDriverWait(driver, 10)
    links = wait.until(EC.presence_of_all_elements_located((By.CSS_SELECTOR, cfg.result_name_links_css)))

    code_n = normalize(code)
    if code_n:
        for a in links:
            if code_n in normalize(a.text or ""):
                return a

    tokens = [normalize(x) for x in (name_hint or "").split() if len(x.strip()) >= 2]
    best, best_score = None, -1
    for a in links:
        t_norm = normalize(a.text or "")
        score = sum(1 for tok in tokens if tok and tok in t_norm)
        if score > best_score:
            best_score = score
            best = a
    return best or (links[0] if links else None)


def open_popup_and_get_fixed_price(driver: webdriver.Chrome, cfg: Config, link_el) -> int | None:
    wait = WebDriverWait(driver, 15)

    main = driver.current_window_handle
    before_count = len(driver.window_handles)

    link_el.click()
    wait.until(lambda d: len(d.window_handles) > before_count)

    new_handle = [h for h in driver.window_handles if h != main][-1]
    driver.switch_to.window(new_handle)

    fixed_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, cfg.fixed_price_css)))
    fixed_val = fixed_input.get_attribute("value")

    driver.close()
    driver.switch_to.window(main)

    return to_int_price(fixed_val)


def process_row(driver: webdriver.Chrome, cfg: Config, ws, row: int):
    code = ""
    if cfg.code_col:
        code = str(ws[f"{cfg.code_col}{row}"].value or "").strip()

    name_hint = str(ws[f"{cfg.name_col}{row}"].value or "").strip()
    excel_price = to_int_price(ws[f"{cfg.excel_price_col}{row}"].value)

    # 검색 우선순위: ✅ 품번 우선 -> 실패 시 상품명
    found = False

    if code:
        if not search_in_goods_list(driver, cfg, code):
            found = True

    if not found and name_hint:
        if not search_in_goods_list(driver, cfg, name_hint):
            found = True

    if not found:
        fill_row_range(ws, row, cfg.row_fill_cols[0], cfg.row_fill_cols[1], FILL_PINK)
        return

    link = pick_product_link(driver, cfg, code=code, name_hint=name_hint)
    if link is None:
        fill_row_range(ws, row, cfg.row_fill_cols[0], cfg.row_fill_cols[1], FILL_PINK)
        return

    mall_fixed = open_popup_and_get_fixed_price(driver, cfg, link)

    # 비교/색칠
    if excel_price is not None and mall_fixed is not None and excel_price == mall_fixed:
        fill_row_range(ws, row, cfg.row_fill_cols[0], cfg.row_fill_cols[1], FILL_GREEN)
        ws[f"{cfg.excel_price_col}{row}"].fill = PatternFill()
    else:
        fill_row_range(ws, row, cfg.row_fill_cols[0], cfg.row_fill_cols[1], None)
        ws[f"{cfg.excel_price_col}{row}"].fill = FILL_YELLOW


# =========================
# main
# =========================
def main():
    # 중지 버튼 팝업 항상 띄우기(별도 스레드)
    t = threading.Thread(target=run_stop_popup, daemon=True)
    t.start()

    # 사용자 입력으로 실행 설정 구성
    cfg = build_cfg_from_user_input()

    wb = load_workbook(cfg.xlsx_path)
    ws = wb[cfg.sheet_name]
    last_row = ws.max_row

    driver = build_driver()
    try:
        wait_user_login(driver, cfg)

        for row in range(cfg.start_row, last_row + 1):
            if stop_flag["stop"]:
                break

            # 행이 완전 공백이면 스킵(중간 빈 줄 대비)
            code_v = ws[f"{cfg.code_col}{row}"].value if cfg.code_col else None
            name_v = ws[f"{cfg.name_col}{row}"].value
            price_v = ws[f"{cfg.excel_price_col}{row}"].value

            if (
                (code_v is None or str(code_v).strip() == "")
                and (name_v is None or str(name_v).strip() == "")
                and (price_v is None or str(price_v).strip() == "")
            ):
                continue

            try:
                process_row(driver, cfg, ws, row)
            except Exception:
                fill_row_range(ws, row, cfg.row_fill_cols[0], cfg.row_fill_cols[1], FILL_PINK)

        out_path = cfg.xlsx_path.replace(".xlsx", "_checked.xlsx")
        wb.save(out_path)
        print(f"완료: {out_path}")

    finally:
        driver.quit()


if __name__ == "__main__":
    main()
